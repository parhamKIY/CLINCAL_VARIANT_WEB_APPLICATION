"""Build the Stage 82 authoritative, editable DOCX report template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates" / "clinical_variant_report_v1.docx"

FONT = "Times New Roman"
BLACK = "000000"
RED = "C00000"
BLUE = "2F5496"
GOLD = "FFC000"
GRAY = "D9D9D9"
GRID = "A5A5A5"
LETTER_WIDTH = Inches(8.5)
LETTER_HEIGHT = Inches(11)
PAGE_MARGIN = Inches(1)
CONTENT_WIDTH = 9360
MAIN_TABLE_WIDTH = 9547


def _set_run_font(run, *, size: float, bold=False, italic=False, color=BLACK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)


def _style(document: Document, name: str, *, size: float, bold=False, italic=False):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    return style


def _configure_styles(document: Document) -> None:
    normal = _style(document, "Normal", size=11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.08

    title = _style(document, "Report Title", size=18, bold=True, italic=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True

    heading = _style(document, "Major Heading", size=12, bold=True)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True

    body = _style(document, "Report Body", size=11)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.08

    metadata = _style(document, "Report Metadata", size=8.5)
    metadata.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    result_variant = _style(document, "Result Variant", size=11, bold=True)
    result_variant.font.color.rgb = RGBColor.from_string(BLUE)
    result_variant.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result_variant.paragraph_format.keep_together = True

    result_class = _style(document, "Result Classification", size=11, bold=True, italic=True)
    result_class.font.color.rgb = RGBColor.from_string(GOLD)
    result_class.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result_class.paragraph_format.keep_together = True

    result_source = _style(document, "Result Source", size=8.5, italic=True)
    result_source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result_source.paragraph_format.keep_together = True

    table = _style(document, "Report Table Text", size=9.5)
    table.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.paragraph_format.keep_together = True

    legend = _style(document, "Report Legend", size=9)
    legend.paragraph_format.space_before = Pt(3)
    legend.paragraph_format.space_after = Pt(3)

    interpretation = _style(document, "Variant Interpretation Body", size=11)
    interpretation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    interpretation.paragraph_format.line_spacing = 1.08
    interpretation.paragraph_format.widow_control = True

    reference = _style(document, "Report Reference", size=7)
    reference.paragraph_format.left_indent = Inches(0.2)
    reference.paragraph_format.first_line_indent = Inches(-0.2)
    reference.paragraph_format.keep_together = True

    notice = _style(document, "Decision Support Notice", size=8.5, italic=True)
    notice.font.color.rgb = RGBColor.from_string("404040")
    notice.paragraph_format.space_before = Pt(8)


def _set_cell_margins(cell, *, top=60, start=80, bottom=60, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_table_borders(table, *, color: str, size: int) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_fixed_table(
    table, widths: list[int], *, caption: str | None = None, indent: int = 80
) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    indent_node = tbl_pr.find(qn("w:tblInd"))
    if indent_node is None:
        indent_node = OxmlElement("w:tblInd")
        tbl_pr.append(indent_node)
    indent_node.set(qn("w:w"), str(indent))
    indent_node.set(qn("w:type"), "dxa")

    if caption:
        caption_node = tbl_pr.find(qn("w:tblCaption"))
        if caption_node is None:
            caption_node = OxmlElement("w:tblCaption")
            tbl_pr.append(caption_node)
        caption_node.set(qn("w:val"), caption)
        description = tbl_pr.find(qn("w:tblDescription"))
        if description is None:
            description = OxmlElement("w:tblDescription")
            tbl_pr.append(description)
        description.set(qn("w:val"), caption)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[min(index, len(widths) - 1)])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_border(cell, edge: str, value: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), value)


def _prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _cell_text(cell, text: str, *, style="Report Table Text", bold=False, underline=False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    _set_run_font(run, size=9.5 if style == "Report Table Text" else 8.5, bold=bold)
    run.underline = underline


def _paragraph(document: Document, text: str, *, style="Report Body", align=None):
    paragraph = document.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    paragraph.add_run(text)
    return paragraph


def _metadata_table(document: Document) -> None:
    table = document.add_table(rows=3, cols=2)
    _set_fixed_table(table, [4680, 4680], caption="Report metadata without patient identifiers")
    table_data = (
        ("Report ID: {{REPORT_ID}}", "Input allele: {{INPUT_INDEX}}"),
        ("Genome build: {{GENOME_BUILD}}", "Report state: {{REPORT_STATE}}"),
        ("Generated: {{GENERATED_DATE}}", "Template: {{TEMPLATE_VERSION}}"),
    )
    for row, values in zip(table.rows, table_data, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            _cell_text(cell, value, style="Report Metadata")
    _set_table_borders(table, color=GRID, size=6)


def _result_block(document: Document) -> None:
    table = document.add_table(rows=1, cols=1)
    _set_fixed_table(
        table,
        [8860],
        caption="Conclusive result for one accepted allele",
        indent=200,
    )
    _set_table_borders(table, color=RED, size=16)
    cell = table.cell(0, 0)
    _set_cell_margins(cell, top=120, start=200, bottom=120, end=200)
    cell.text = ""
    variant = cell.paragraphs[0]
    variant.style = "Result Variant"
    variant.add_run("{{VARIANT_DISPLAY}}")
    zygosity = cell.add_paragraph(style="Result Variant")
    zygosity.add_run("{{ZYGOSITY}}")
    classification = cell.add_paragraph(style="Result Classification")
    classification.add_run("{{CLASSIFICATION}}")
    source = cell.add_paragraph(style="Result Source")
    source.add_run("Source: {{CLASSIFICATION_SOURCE}}")


def _main_findings_tables(document: Document) -> None:
    widths_a = [1518, 1623, 1079, 993, 821, 2253, 1260]
    table_a = document.add_table(rows=3, cols=7)
    _set_fixed_table(table_a, widths_a, caption="{{MAIN_FINDINGS_TABLE}}")
    _set_table_borders(table_a, color=GRID, size=7)

    gene = table_a.cell(0, 0)
    variant = table_a.cell(0, 1)
    population = table_a.cell(0, 3)
    disease = table_a.cell(0, 5)
    zygosity = table_a.cell(0, 6)
    for cell, text in (
        (gene, "Gene & transcript"),
        (variant, "Variant"),
        (population, "Population frequency"),
        (disease, "Associated disease & ID"),
        (zygosity, "Zygosity"),
    ):
        _cell_text(cell, text, bold=True)
    for column in (0, 1, 5, 6):
        _set_cell_border(table_a.cell(0, column), "bottom", "nil")
        _set_cell_border(table_a.cell(1, column), "top", "nil")
    _set_cell_border(table_a.cell(0, 2), "right", "nil")
    _set_cell_border(table_a.cell(0, 3), "left", "nil")
    _set_cell_border(table_a.cell(0, 3), "right", "nil")
    _set_cell_border(table_a.cell(0, 4), "left", "nil")
    for index, text in enumerate(("{{POP_SOURCE_1}}", "{{POP_SOURCE_2}}", "{{POP_SOURCE_3}}"), start=2):
        _cell_text(table_a.cell(1, index), text, bold=True)
    for cell, text in zip(
        table_a.rows[2].cells,
        (
            "{{GENE_TRANSCRIPT}}",
            "{{HGVS_DISPLAY}}",
            "{{POP_VALUE_1}}",
            "{{POP_VALUE_2}}",
            "{{POP_VALUE_3}}",
            "{{DISEASE_ASSOCIATION}}",
            "{{ZYGOSITY}}",
        ),
        strict=True,
    ):
        _cell_text(cell, text)
        _shade(cell, GRAY)
    _repeat_header(table_a.rows[0])
    _repeat_header(table_a.rows[1])
    _prevent_split(table_a.rows[2])

    widths_b = [1518, 1623, 1528, 1346, 2253, 1279]
    table_b = document.add_table(rows=2, cols=6)
    _set_fixed_table(table_b, widths_b, caption="Main findings evidence continuation")
    _set_table_borders(table_b, color=GRID, size=7)
    headers = (
        "{{COMPUTATIONAL_SOURCE_1}}",
        "{{COMPUTATIONAL_SOURCE_2}}",
        "{{COMPUTATIONAL_SOURCE_3}}",
        "Stable variant ID",
        "Classification context",
        "Inheritance",
    )
    values = (
        "{{COMPUTATIONAL_VALUE_1}}",
        "{{COMPUTATIONAL_VALUE_2}}",
        "{{COMPUTATIONAL_VALUE_3}}",
        "{{STABLE_VARIANT_ID}}",
        "{{CLASSIFICATION_CONTEXT}}",
        "{{INHERITANCE}}",
    )
    for cell, value in zip(table_b.rows[0].cells, headers, strict=True):
        _cell_text(cell, value, bold=True)
    for cell, value in zip(table_b.rows[1].cells, values, strict=True):
        _cell_text(cell, value)
        _shade(cell, GRAY)
    _repeat_header(table_b.rows[0])
    _prevent_split(table_b.rows[1])


def _data_sources_table(document: Document) -> None:
    widths = [1440, 1440, 1080, 2160, 1800, 1440]
    table = document.add_table(rows=2, cols=6)
    _set_fixed_table(table, widths, caption="{{DATA_SOURCES}}")
    _set_table_borders(table, color=GRID, size=7)
    headers = ("Source", "Capability", "Status", "Record or dataset", "Method", "Link")
    values = (
        "{{DATA_SOURCE_NAME}}",
        "{{DATA_SOURCE_CAPABILITY}}",
        "{{DATA_SOURCE_STATUS}}",
        "{{DATA_SOURCE_RECORD}}",
        "{{DATA_SOURCE_METHOD}}",
        "{{DATA_SOURCE_LINK}}",
    )
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        _cell_text(cell, value, bold=True)
    for cell, value in zip(table.rows[1].cells, values, strict=True):
        _cell_text(cell, value)
        _shade(cell, GRAY)
    _repeat_header(table.rows[0])
    _prevent_split(table.rows[1])


def build_template(output: Path = OUTPUT) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = LETTER_WIDTH
    section.page_height = LETTER_HEIGHT
    section.top_margin = PAGE_MARGIN
    section.right_margin = PAGE_MARGIN
    section.bottom_margin = PAGE_MARGIN
    section.left_margin = PAGE_MARGIN
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    _configure_styles(document)
    document.core_properties.title = "Clinical Variant Report Template V1"
    document.core_properties.subject = "One editable report per accepted allele"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.keywords = ""
    document.core_properties.comments = ""

    _paragraph(document, "{{REPORT_TITLE}}", style="Report Title")
    _metadata_table(document)
    _paragraph(document, "Clinical Features:", style="Major Heading")
    _paragraph(document, "{{CLINICAL_FEATURES}}")
    _paragraph(document, "Method:", style="Major Heading")
    _paragraph(document, "{{METHOD_SCOPE}}")
    _paragraph(document, "Conclusive Result(s):", style="Major Heading")
    _result_block(document)
    _paragraph(document, "Brief Interpretation(s):", style="Major Heading")
    _paragraph(document, "{{BRIEF_INTERPRETATION}}")
    _paragraph(document, "{{DECISION_SUPPORT_NOTICE}}", style="Decision Support Notice")

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    _paragraph(document, "Main Finding(s) in Detail:", style="Major Heading")
    _main_findings_tables(document)
    _paragraph(document, "{{TABLE_LEGENDS}}", style="Report Legend")
    _paragraph(document, "Variant interpretation:", style="Major Heading")
    allele = _paragraph(document, "", style="Report Body")
    allele.paragraph_format.keep_with_next = True
    allele_run = allele.add_run("{{VARIANT_DISPLAY}}")
    allele_run.underline = True
    _paragraph(document, "{{VARIANT_INTERPRETATION}}", style="Variant Interpretation Body")

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    _paragraph(document, "Variant(s) classification:", style="Major Heading")
    _paragraph(document, "{{CLASSIFICATION_SUMMARY}}")
    _paragraph(document, "Method:", style="Major Heading")
    _paragraph(document, "{{METHOD_DETAILS}}")
    _paragraph(document, "Comments / scope:", style="Major Heading")
    _paragraph(document, "{{COMMENTS_AND_SCOPE}}")
    _paragraph(document, "References:", style="Major Heading")
    _paragraph(document, "{{REFERENCES}}", style="Report Reference")
    _paragraph(document, "Data Sources:", style="Major Heading")
    _data_sources_table(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


if __name__ == "__main__":
    print(build_template())
