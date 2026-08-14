"""Stage 101 structural and optional raster regression harness for DOCX reports."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.references import canonicalize_reference
from backend.report_data import ReportData, validate_report_data
from backend.report_docx import render_report_data_docx


SCENARIO_FIXTURE = ROOT / "tests" / "fixtures" / "stage101_visual_scenarios.json"
BASELINE_PATH = ROOT / "tests" / "golden" / "stage101" / "snapshots.json"

EXPECTED_HEADINGS = [
    "NGS Result Report",
    "Clinical Features:",
    "Method:",
    "Variant and source classification context:",
    "Brief Interpretation(s):",
    "Main Finding(s) in Detail:",
    "Variant interpretation:",
    "Variant(s) classification:",
    "Method:",
    "Comments / scope:",
    "References:",
    "Data Sources:",
]
REQUIRED_LABELS = frozenset(
    {
        "NGS Result Report",
        "Clinical Features:",
        "Variant and source classification context:",
        "Brief Interpretation(s):",
        "Main Finding(s) in Detail:",
        "Variant interpretation:",
        "Variant(s) classification:",
        "References:",
        "Data Sources:",
        "Gene & transcript",
        "Variant",
        "Inheritance",
    }
)
_PAGE_NUMBER_PATTERN = re.compile(r"page-(\d+)\.png$", re.I)


class VisualRegressionError(AssertionError):
    """Raised when a report violates the Stage 101 layout contract."""


def load_scenario_registry() -> dict[str, object]:
    """Load the deterministic Stage 101 scenario registry."""

    return json.loads(SCENARIO_FIXTURE.read_text(encoding="utf-8"))


def _scenario_base(index: int, name: str) -> dict[str, Any]:
    registry = load_scenario_registry()
    base_path = SCENARIO_FIXTURE.parent / str(registry["base_fixture"])
    report = json.loads(base_path.read_text(encoding="utf-8"))
    report["report_id"] = f"stage101-{name}-report"
    report["analysis_id"] = "stage101-visual-regression"
    report["input_index"] = index
    return report


def _replace_interpretation(report: dict[str, Any], text: str) -> None:
    interpretation = report["interpretation"]
    interpretation["original_model_interpretation"] = text
    interpretation["current_reviewer_interpretation"] = text
    interpretation["edit_history"] = []


def _sparse_evidence(report: dict[str, Any]) -> None:
    report["phenotype_summary"] = {
        "accepted_hpo_terms": [],
        "matched_hpo_terms": [],
        "concordance": "unavailable",
        "score": None,
        "summary": "Phenotype evidence is unavailable for this synthetic case.",
        "evidence_status": "unavailable",
    }
    report["conclusive_result"].update(
        {
            "classification": None,
            "classification_source": None,
            "status": "not_assessed",
        }
    )
    report["main_findings"] = {
        "population_frequencies": [],
        "disease_associations": [],
        "computational_evidence": [],
        "stable_variant_ids": [],
        "classifications": [],
    }
    text = (
        "Evidence is sparse for this synthetic allele. No exact population, "
        "disease, computational, stable-identifier, or classification record "
        "is available. The report remains reviewable and does not infer missing "
        "evidence."
    )
    _replace_interpretation(report, text)
    report["classification_summary"].update(
        {
            "reviewer_confirmed_classification": None,
            "clinvar_classification": None,
            "automated_classification": None,
            "conflict_status": "unavailable",
            "conflict_severity": None,
            "source_attributions": [],
            "summary": "Classification evidence is unavailable.",
        }
    )
    report["literature_references"] = []
    report["data_sources"] = []
    report["warnings"] = [
        {
            "severity": "PARTIAL",
            "code": "sparse_synthetic_evidence",
            "message": "Most evidence capabilities are unavailable.",
            "capability": None,
        }
    ]
    report["review_state"]["reviewer_summary"] = text
    report["review_state"]["reviewer_notes"] = []


def _long_interpretation(report: dict[str, Any]) -> None:
    source = report["interpretation"]["current_reviewer_interpretation"]
    text = "\n\n".join(
        f"Layout stress paragraph {index + 1}. {source}" for index in range(8)
    )
    _replace_interpretation(report, text)
    report["review_state"]["reviewer_summary"] = (
        "This synthetic case verifies that a long reviewer-approved interpretation "
        "flows across pages without fixed-height clipping or silent truncation."
    )


def _many_references(report: dict[str, Any]) -> None:
    report["literature_references"] = [
        canonicalize_reference(
            source="DOI",
            identifier=f"10.1000/stage101.{index:02d}",
            title=(
                "Synthetic visual-regression reference with deliberately descriptive "
                f"bibliographic text {index:02d}"
            ),
            reference_id=f"R{index}",
        )
        for index in range(1, 21)
    ]


def _no_phenotype_match(report: dict[str, Any]) -> None:
    phenotype = report["phenotype_summary"]
    phenotype["matched_hpo_terms"] = []
    phenotype["concordance"] = "no_supported_association"
    phenotype["score"] = 0.0
    phenotype["summary"] = (
        "No supported association was retained between the accepted phenotype and "
        "this synthetic allele."
    )
    phenotype["evidence_status"] = "no_match"
    text = (
        "The allele remains interpretable, but the retained evidence does not support "
        "an association with the accepted phenotype. Variant-level population, "
        "classification, and computational context remains source attributed and "
        "must not be converted into a phenotype match."
    )
    _replace_interpretation(report, text)
    report["review_state"]["reviewer_summary"] = text


def _complex_indel(report: dict[str, Any]) -> None:
    reference = "ACGT" * 38
    alternate = "A"
    hgvs_c = "c.101_252delinsA"
    hgvs_p = "p.(Gly34_Leu84delinsAsp)"
    report["variant_identity"].update(
        {
            "position": 200_001,
            "reference": reference,
            "alternate": alternate,
            "hgvs_c": hgvs_c,
            "hgvs_p": hgvs_p,
        }
    )
    report["conclusive_result"].update(
        {"hgvs_c": hgvs_c, "hgvs_p": hgvs_p}
    )
    report["review_state"]["reviewer_summary"] = (
        "This synthetic complex deletion-insertion verifies wrapping of long allele "
        "identity without merging, clipping, or loss of the exact reference allele."
    )


_MUTATIONS = {
    "fully_populated": lambda report: None,
    "sparse_evidence": _sparse_evidence,
    "long_interpretation": _long_interpretation,
    "many_references": _many_references,
    "no_phenotype_match": _no_phenotype_match,
    "complex_indel": _complex_indel,
}


def build_scenario_report(name: str) -> ReportData:
    """Build one validated synthetic ReportData fixture by registered scenario."""

    registry = load_scenario_registry()
    names = [item["name"] for item in registry["scenarios"]]
    if name not in names or name not in _MUTATIONS:
        raise VisualRegressionError(f"Unknown Stage 101 scenario: {name}")
    report = _scenario_base(names.index(name), name)
    _MUTATIONS[name](report)
    return validate_report_data(report)


def scenario_page_range(name: str) -> tuple[int, int]:
    registry = load_scenario_registry()
    for scenario in registry["scenarios"]:
        if scenario["name"] == name:
            minimum, maximum = scenario["expected_page_range"]
            return int(minimum), int(maximum)
    raise VisualRegressionError(f"Unknown Stage 101 scenario: {name}")


def _table_caption(table: object) -> str | None:
    node = table._tbl.tblPr.find(qn("w:tblCaption"))
    return node.get(qn("w:val")) if node is not None else None


def _page_segments(document_xml: bytes) -> list[str]:
    root = ElementTree.fromstring(document_xml)
    segments: list[list[str]] = [[]]
    for element in root.iter():
        if element.tag == qn("w:br") and element.get(qn("w:type")) == "page":
            segments.append([])
        elif element.tag == qn("w:t") and element.text:
            segments[-1].append(element.text)
    return [" ".join(segment).strip() for segment in segments]


def inspect_docx(data: bytes) -> dict[str, object]:
    """Return a deterministic non-text-only layout snapshot for one DOCX."""

    document = Document(BytesIO(data))
    with ZipFile(BytesIO(data)) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Report Title", "Major Heading"}
    ]
    all_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    tables = [
        {
            "caption": _table_caption(table),
            "rows": len(table.rows),
            "columns": len(table.columns),
        }
        for table in document.tables
    ]
    segments = _page_segments(document_xml)
    exact_height_rows = sum(
        row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY
        for table in document.tables
        for row in table.rows
    )
    section = document.sections[0]
    snapshot: dict[str, object] = {
        "docx_sha256": hashlib.sha256(data).hexdigest(),
        "styles_sha256": hashlib.sha256(styles_xml).hexdigest(),
        "page_width_twips": int(section.page_width.twips),
        "page_height_twips": int(section.page_height.twips),
        "margins_twips": [
            int(section.top_margin.twips),
            int(section.right_margin.twips),
            int(section.bottom_margin.twips),
            int(section.left_margin.twips),
        ],
        "explicit_page_breaks": len(segments) - 1,
        "blank_explicit_page_segments": sum(not segment for segment in segments),
        "headings": headings,
        "missing_required_labels": sorted(
            label for label in REQUIRED_LABELS if label not in all_text
        ),
        "tables": tables,
        "exact_height_rows": exact_height_rows,
        "unresolved_placeholders": "{{" in all_text or b"{{" in document_xml,
        "hyperlink_count": document_xml.count(b"w:hyperlink"),
    }
    signature_payload = {
        key: value for key, value in snapshot.items() if key != "docx_sha256"
    }
    snapshot["layout_signature"] = hashlib.sha256(
        json.dumps(
            signature_payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    return snapshot


def assert_structural_contract(snapshot: dict[str, object]) -> None:
    """Reject silent damage to the professor-family Word layout."""

    if snapshot["headings"] != EXPECTED_HEADINGS:
        raise VisualRegressionError("Clinical report heading order changed.")
    if snapshot["missing_required_labels"]:
        raise VisualRegressionError("Required clinical report labels are missing.")
    if len(snapshot["tables"]) != 5:
        raise VisualRegressionError("Clinical report table structure changed.")
    captions = {
        table["caption"] for table in snapshot["tables"] if table["caption"]
    }
    if not {"Main Findings", "Data Sources"}.issubset(captions):
        raise VisualRegressionError("Required report tables are unavailable.")
    if snapshot["explicit_page_breaks"] != 2:
        raise VisualRegressionError("Controlled report page breaks changed.")
    if snapshot["blank_explicit_page_segments"]:
        raise VisualRegressionError("The DOCX contains an empty controlled page.")
    if snapshot["exact_height_rows"]:
        raise VisualRegressionError("Fixed-height rows may clip report content.")
    if snapshot["unresolved_placeholders"]:
        raise VisualRegressionError("The DOCX contains unresolved placeholders.")
    if snapshot["page_width_twips"] != 12_240 or snapshot[
        "page_height_twips"
    ] != 15_840:
        raise VisualRegressionError("Report page geometry changed.")
    if snapshot["margins_twips"] != [1_440, 1_440, 1_440, 1_440]:
        raise VisualRegressionError("Report margins changed.")


def baseline_projection(snapshot: dict[str, object]) -> dict[str, object]:
    """Keep the committed binary/visual signature compact and reviewable."""

    return {
        key: snapshot[key]
        for key in (
            "docx_sha256",
            "styles_sha256",
            "layout_signature",
            "tables",
            "hyperlink_count",
        )
    }


def _average_hash(path: Path) -> str:
    from PIL import Image

    with Image.open(path) as image:
        grayscale = image.convert("L").resize((8, 8))
        pixels = list(grayscale.get_flattened_data())
    average = sum(pixels) / len(pixels)
    bits = "".join("1" if pixel < average else "0" for pixel in pixels)
    return f"{int(bits, 2):016x}"


def analyze_rendered_pages(
    paths: list[Path],
    *,
    expected_page_range: tuple[int, int],
) -> dict[str, object]:
    """Check raster page count, blank pages, edge overflow, and image hashes."""

    from PIL import Image, ImageStat

    ordered = sorted(
        paths,
        key=lambda path: int(_PAGE_NUMBER_PATTERN.search(path.name).group(1)),
    )
    minimum, maximum = expected_page_range
    if not minimum <= len(ordered) <= maximum:
        raise VisualRegressionError("Rendered report page count is out of range.")
    hashes: list[str] = []
    for path in ordered:
        with Image.open(path) as image:
            grayscale = image.convert("L")
            inverted = grayscale.point(lambda pixel: 255 - pixel)
            if ImageStat.Stat(inverted).mean[0] < 0.4:
                raise VisualRegressionError(f"Unexpected blank page: {path.name}")
            width, height = grayscale.size
            edge = max(2, min(width, height) // 300)
            borders = (
                inverted.crop((0, 0, width, edge)),
                inverted.crop((0, height - edge, width, height)),
                inverted.crop((0, 0, edge, height)),
                inverted.crop((width - edge, 0, width, height)),
            )
            if any(ImageStat.Stat(border).mean[0] > 1.0 for border in borders):
                raise VisualRegressionError(
                    f"Rendered content touches a page edge: {path.name}"
                )
        hashes.append(_average_hash(path))
    return {"page_count": len(ordered), "average_hashes": hashes}


def compare_image_hashes(
    actual: list[str],
    expected: list[str],
    *,
    tolerance: int = 6,
) -> None:
    """Compare average hashes with a small renderer-tolerant Hamming distance."""

    if len(actual) != len(expected):
        raise VisualRegressionError("Rendered snapshot page count changed.")
    for page, (left, right) in enumerate(zip(actual, expected, strict=True), 1):
        distance = (int(left, 16) ^ int(right, 16)).bit_count()
        if distance > tolerance:
            raise VisualRegressionError(
                f"Rendered page {page} exceeded image tolerance."
            )


def _render_preview(
    docx_path: Path,
    output_dir: Path,
    *,
    renderer_python: Path,
    renderer_script: Path,
) -> list[Path]:
    command = [
        str(renderer_python),
        str(renderer_script),
        str(docx_path),
        "--output_dir",
        str(output_dir),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    if result.returncode:
        detail = (result.stderr or result.stdout).strip().splitlines()
        raise VisualRegressionError(
            detail[-1] if detail else "DOCX preview rendering failed."
        )
    return list(output_dir.glob("page-*.png"))


def _native_command(executable: str) -> list[str]:
    if os.name == "nt" and Path(executable).suffix.casefold() in {".bat", ".cmd"}:
        return ["cmd", "/c", executable]
    return [executable]


def render_docx_with_local_office(data: bytes, output_dir: Path) -> list[Path]:
    """Render a DOCX through installed LibreOffice and Poppler executables."""

    soffice = shutil.which("soffice")
    pdftoppm = shutil.which("pdftoppm")
    if soffice is None or pdftoppm is None:
        raise VisualRegressionError(
            "LibreOffice and Poppler are required for real raster checks."
        )
    output_dir.mkdir(parents=True, exist_ok=True)
    source = output_dir / "report.docx"
    source.write_bytes(data)
    profile = output_dir / "libreoffice_profile"
    profile.mkdir(exist_ok=True)
    conversion = subprocess.run(
        _native_command(soffice)
        + [
            "--headless",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )
    pdf = output_dir / "report.pdf"
    if conversion.returncode or not pdf.is_file():
        raise VisualRegressionError("LibreOffice DOCX rendering failed.")
    raster = subprocess.run(
        _native_command(pdftoppm)
        + ["-png", "-r", "120", str(pdf), str(output_dir / "page")],
        capture_output=True,
        text=True,
        timeout=120,
    )
    pages = list(output_dir.glob("page-*.png"))
    if raster.returncode or not pages:
        raise VisualRegressionError("Rendered PDF rasterization failed.")
    return pages


def run_harness(
    output_dir: Path,
    *,
    renderer_python: Path | None = None,
    renderer_script: Path | None = None,
    use_local_office: bool = False,
) -> dict[str, object]:
    """Run structural snapshots and optional real DOCX raster checks."""

    registry = load_scenario_registry()
    baseline = json.loads(BASELINE_PATH.read_text(encoding="utf-8"))
    results: dict[str, object] = {}
    output_dir.mkdir(parents=True, exist_ok=True)
    for scenario in registry["scenarios"]:
        name = scenario["name"]
        data = render_report_data_docx(build_scenario_report(name))
        snapshot = inspect_docx(data)
        assert_structural_contract(snapshot)
        if baseline_projection(snapshot) != baseline["scenarios"][name]:
            raise VisualRegressionError(f"Structural snapshot changed: {name}")
        docx_path = output_dir / f"{name}.docx"
        docx_path.write_bytes(data)
        result: dict[str, object] = {"structural": snapshot}
        if use_local_office:
            pages = render_docx_with_local_office(data, output_dir / f"{name}_preview")
            result["rendered"] = analyze_rendered_pages(
                pages,
                expected_page_range=scenario_page_range(name),
            )
        elif renderer_python is not None and renderer_script is not None:
            preview_dir = output_dir / f"{name}_preview"
            pages = _render_preview(
                docx_path,
                preview_dir,
                renderer_python=renderer_python,
                renderer_script=renderer_script,
            )
            result["rendered"] = analyze_rendered_pages(
                pages,
                expected_page_range=scenario_page_range(name),
            )
        results[name] = result
    return {"schema_version": "1.0", "scenarios": results}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--renderer-python", type=Path)
    parser.add_argument("--renderer-script", type=Path)
    parser.add_argument("--local-office", action="store_true")
    args = parser.parse_args()
    if (args.renderer_python is None) != (args.renderer_script is None):
        parser.error("Both renderer arguments are required for raster checks.")
    if args.local_office and args.renderer_python is not None:
        parser.error("Choose either --local-office or explicit renderer arguments.")
    try:
        result = run_harness(
            args.output_dir,
            renderer_python=args.renderer_python,
            renderer_script=args.renderer_script,
            use_local_office=args.local_office,
        )
    except (OSError, VisualRegressionError, ValueError) as exc:
        print(f"Stage 101 visual regression: FAILED: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
