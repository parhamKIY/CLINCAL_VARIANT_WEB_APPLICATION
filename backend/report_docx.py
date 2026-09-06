"""Professor-template DOCX rendering for one validated ReportData V4 allele."""

from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from backend.references import validated_reference_url
from backend.report_data import ReportData, ReportDataError, validate_report_data
from backend.report_narrative import substantive_interpretation_narrative


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE_PATH = ROOT / "templates" / "clinical_variant_report_v1.docx"
SUPPORTED_TEMPLATE_VERSION = "professor-report-v1"
MAX_DOCX_BYTES = 5 * 1024 * 1024
PLACEHOLDER_PATTERN = re.compile(r"\{\{[A-Z0-9_]+\}\}")

REQUIRED_PLACEHOLDERS = frozenset(
    {
        "{{REPORT_TITLE}}",
        "{{CLINICAL_FEATURES}}",
        "{{VARIANT_DISPLAY}}",
        "{{ZYGOSITY}}",
        "{{CLASSIFICATION}}",
        "{{BRIEF_INTERPRETATION}}",
        "{{MAIN_FINDINGS_TABLE}}",
        "{{VARIANT_INTERPRETATION}}",
        "{{CLASSIFICATION_SUMMARY}}",
        "{{REFERENCES}}",
        "{{DATA_SOURCES}}",
    }
)


class ReportDocxError(ValueError):
    """Raised when a validated allele report cannot be rendered safely."""


def report_docx_filename(value: object) -> str:
    """Return the stable one-based artifact name for one ReportData record."""

    report = validate_report_data(value)
    return f"variant_{report['input_index'] + 1:03d}_report.docx"


def _status_text(status: str, *, no_match: str = "No exact record") -> str:
    labels = {
        "available": "Available",
        "no_match": no_match,
        "unsupported": "Not supported by this source",
        "unavailable": "Not available",
        "not_assessed": "Not assessed",
    }
    return labels[status]


def _sentence_case(value: str | None, fallback: str = "Not available") -> str:
    if not value:
        return fallback
    cleaned = value.replace("_", " ").strip()
    if cleaned.lower() in {
        "vus",
        "uncertain significance",
        "variant of uncertain significance",
    }:
        return "Uncertain significance (VUS)"
    return cleaned[:1].upper() + cleaned[1:]


def _variant_display(report: ReportData) -> str:
    variant = report["variant_identity"]
    gene = variant["gene"] or "Gene not available"
    parts = [value for value in (variant["hgvs_c"], variant["hgvs_p"]) if value]
    if parts:
        return f"{gene}:{' '.join(parts)}"
    return (
        f"{gene}: {variant['genome_build']} chr{variant['chromosome']}:"
        f"{variant['position']} {variant['reference']}>{variant['alternate']}"
    )


def _gene_transcript(report: ReportData) -> str:
    variant = report["variant_identity"]
    values = [value for value in (variant["gene"], variant["transcript"]) if value]
    return "\n".join(values) if values else "Not available"


def _hgvs_display(report: ReportData) -> str:
    variant = report["variant_identity"]
    values = [value for value in (variant["hgvs_c"], variant["hgvs_p"]) if value]
    if values:
        return "\n".join(values)
    return (
        f"{variant['genome_build']} chr{variant['chromosome']}:"
        f"{variant['position']} {variant['reference']}>{variant['alternate']}"
    )


def _phenotype_text(report: ReportData) -> str:
    phenotype = report["phenotype_summary"]
    terms = [
        f"{item['label']} ({item['hpo_id']})" if item["label"] else item["hpo_id"]
        for item in phenotype["accepted_hpo_terms"]
    ]
    term_text = "; ".join(terms) if terms else "No accepted HPO terms available"
    concordance = phenotype["concordance"].replace("_", " ")
    summary = phenotype["summary"] or _status_text(
        phenotype["evidence_status"], no_match="No supported association"
    )
    return (
        f"Accepted features: {term_text}. Phenotype concordance: "
        f"{concordance}. {summary}"
    )


def _brief_interpretation(report: ReportData) -> str:
    preferred = report["review_state"]["reviewer_summary"]
    if preferred:
        return preferred
    interpretation = report["interpretation"]["current_reviewer_interpretation"]
    if not interpretation:
        return "Interpretation is not available and requires human review."
    return substantive_interpretation_narrative(interpretation)


def _frequency_text(item: dict[str, object]) -> str:
    if item["status"] != "available":
        return _status_text(str(item["status"]))
    frequency = float(item["frequency"])
    if frequency != 0 and frequency < 0.001:
        return f"{frequency:.3e}"
    return f"{frequency:.6g}"


def _population_slots(report: ReportData) -> tuple[list[str], list[str]]:
    findings = report["main_findings"]["population_frequencies"][:3]
    names: list[str] = []
    values: list[str] = []
    for item in findings:
        label = item["source"]
        if item["dataset"]:
            label = f"{label}\n{item['dataset']}"
        if item["fallback_used"]:
            label = f"{label}\n(fallback)"
        names.append(label)
        values.append(_frequency_text(dict(item)))
    while len(names) < 3:
        names.append("Not available")
        values.append("Not available")
    return names, values


def _disease_text(report: ReportData) -> tuple[str, str]:
    findings = report["main_findings"]["disease_associations"]
    if not findings:
        return "Not available", "Not available"
    item = findings[0]
    if item["status"] != "available":
        missing = _status_text(
            item["status"], no_match="No validated association found"
        )
        return missing, "Not available"
    identity = item["disease_name"] or "Disease name not available"
    if item["disease_id"]:
        identity = f"{identity}\n({item['disease_id']})"
    return identity, item["inheritance"] or "Not available"


def _computational_slots(report: ReportData) -> tuple[list[str], list[str]]:
    findings = report["main_findings"]["computational_evidence"][:3]
    names: list[str] = []
    values: list[str] = []
    for item in findings:
        names.append(item["tool"] or item["source"])
        if item["status"] != "available":
            values.append(_status_text(item["status"]))
            continue
        parts = [item["prediction"]] if item["prediction"] else []
        if item["score"] is not None:
            parts.append(f"score {float(item['score']):.6g}")
        values.append("\n".join(parts))
    while len(names) < 3:
        names.append("Not available")
        values.append("Not available")
    return names, values


def _stable_identifiers(report: ReportData) -> str:
    values: list[str] = []
    for item in report["main_findings"]["stable_variant_ids"]:
        if item["status"] == "available" and item["identifier"]:
            values.append(f"{item['identifier_type']}: {item['identifier']}")
    return "\n".join(values) if values else "Not available"


def _classification_context(report: ReportData) -> str:
    values: list[str] = []
    for item in report["main_findings"]["classifications"][:3]:
        if item["status"] == "available" and item["classification"]:
            text = f"{item['source']}: {_sentence_case(item['classification'])}"
            if item["review_status"]:
                text += f" ({item['review_status']})"
            values.append(text)
        else:
            values.append(f"{item['source']}: {_status_text(item['status'])}")
    return "\n".join(values) if values else "Not available"


def _classification_summary(report: ReportData) -> str:
    summary = report["classification_summary"]
    statements: list[str] = []
    if summary["summary"]:
        statements.append(summary["summary"])
    for label, value in (
        ("Reviewer-confirmed classification", summary["reviewer_confirmed_classification"]),
        ("ClinVar classification", summary["clinvar_classification"]),
        ("Automated classification context", summary["automated_classification"]),
    ):
        statements.append(f"{label}: {_sentence_case(value)}.")
    statements.append(
        f"Conflict status: {summary['conflict_status'].replace('_', ' ')}."
    )
    if summary["source_attributions"]:
        statements.append(
            "Sources: " + ", ".join(summary["source_attributions"]) + "."
        )
    statements.append(
        "The application has not independently adjudicated an ACMG/AMP classification."
    )
    return " ".join(statements)


def _comments_and_scope(report: ReportData) -> str:
    comments = [item["message"] for item in report["warnings"]]
    comments.extend(report["review_state"]["reviewer_notes"])
    comments.append(
        "This allele-level report supports qualified human review and does not provide "
        "a diagnosis, treatment recommendation, or testing directive."
    )
    return "\n".join(comments)


def _placeholder_map(report: ReportData) -> dict[str, str]:
    populations, population_values = _population_slots(report)
    computational, computational_values = _computational_slots(report)
    disease, inheritance = _disease_text(report)
    conclusive = report["conclusive_result"]
    classification_summary = report["classification_summary"]
    independent_classification = (
        classification_summary["independent_acmg_adjudication"]
        or classification_summary["reviewer_confirmed_classification"] is not None
    )
    draft_ai_classification = (
        conclusive["classification_source"] == "LLM draft classification"
    )
    display_classification = independent_classification or draft_ai_classification
    classification = (
        _sentence_case(conclusive["classification"])
        if display_classification and conclusive["status"] == "available"
        else "Not independently determined"
    )
    classification_source = (
        conclusive["classification_source"]
        if display_classification and conclusive["classification_source"]
        else "No independent application classification"
    )
    interpretation = report["interpretation"]["current_reviewer_interpretation"]
    if not interpretation:
        interpretation = "Interpretation is not available and requires human review."
    generated_date = report["provenance"]["generated_at"].split("T", 1)[0]
    return {
        "{{REPORT_TITLE}}": "NGS Result Report",
        "{{REPORT_ID}}": report["report_id"],
        "{{INPUT_INDEX}}": str(report["input_index"] + 1),
        "{{GENOME_BUILD}}": report["variant_identity"]["genome_build"],
        "{{REPORT_STATE}}": report["review_state"]["review_status"].capitalize(),
        "{{GENERATED_DATE}}": generated_date,
        "{{TEMPLATE_VERSION}}": report["template_version"],
        "{{CLINICAL_FEATURES}}": _phenotype_text(report),
        "{{METHOD_SCOPE}}": (
            "Allele-level evidence synthesis for an already filtered variant; this "
            "application did not perform sequencing or genome-wide prioritization."
        ),
        "{{VARIANT_DISPLAY}}": _variant_display(report),
        "{{ZYGOSITY}}": _sentence_case(conclusive["zygosity"]),
        "{{CLASSIFICATION}}": classification,
        "{{CLASSIFICATION_SOURCE}}": classification_source,
        "{{BRIEF_INTERPRETATION}}": _brief_interpretation(report),
        "{{DECISION_SUPPORT_NOTICE}}": (
            "Decision-support output for educational and research use; qualified "
            "human review is required."
        ),
        "{{GENE_TRANSCRIPT}}": _gene_transcript(report),
        "{{HGVS_DISPLAY}}": _hgvs_display(report),
        "{{POP_SOURCE_1}}": populations[0],
        "{{POP_SOURCE_2}}": populations[1],
        "{{POP_SOURCE_3}}": populations[2],
        "{{POP_VALUE_1}}": population_values[0],
        "{{POP_VALUE_2}}": population_values[1],
        "{{POP_VALUE_3}}": population_values[2],
        "{{DISEASE_ASSOCIATION}}": disease,
        "{{COMPUTATIONAL_SOURCE_1}}": computational[0],
        "{{COMPUTATIONAL_SOURCE_2}}": computational[1],
        "{{COMPUTATIONAL_SOURCE_3}}": computational[2],
        "{{COMPUTATIONAL_VALUE_1}}": computational_values[0],
        "{{COMPUTATIONAL_VALUE_2}}": computational_values[1],
        "{{COMPUTATIONAL_VALUE_3}}": computational_values[2],
        "{{STABLE_VARIANT_ID}}": _stable_identifiers(report),
        "{{CLASSIFICATION_CONTEXT}}": _classification_context(report),
        "{{INHERITANCE}}": inheritance,
        "{{TABLE_LEGENDS}}": (
            "Population values are exact-allele observations from the named source. "
            "Computational predictions are source-attributed and are not independent "
            "ACMG classifications. Missingness is written explicitly."
        ),
        "{{VARIANT_INTERPRETATION}}": interpretation,
        "{{CLASSIFICATION_SUMMARY}}": _classification_summary(report),
        "{{METHOD_DETAILS}}": (
            "The application validated the supplied allele identity, collected bounded "
            "source-attributed evidence, preserved provider availability and fallback "
            "provenance, and rendered the reviewer-approved interpretation without a "
            "new model request."
        ),
        "{{COMMENTS_AND_SCOPE}}": _comments_and_scope(report),
    }


def _all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_paragraph_text(paragraph, replacements: dict[str, str]) -> None:
    runs = list(paragraph.runs)
    for run, original_text in zip(
        runs,
        (run.text for run in runs),
        strict=True,
    ):
        updated_text = original_text
        for token, value in replacements.items():
            if token in updated_text:
                updated_text = updated_text.replace(token, value)
        if updated_text != original_text:
            run.text = updated_text
    paragraph_text = paragraph.text
    remaining = [token for token in replacements if token in paragraph_text]
    if not remaining:
        return
    text = paragraph_text
    for token in remaining:
        text = text.replace(token, replacements[token])
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_paragraph_text(paragraph, text: str, *, size: float | None = None) -> None:
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().rFonts
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")


def _add_hyperlink(paragraph, label: str, url: str, *, size_half_points: int) -> None:
    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(size_half_points))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5496")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((fonts, size, color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _find_paragraph(document: Document, text: str):
    for paragraph in _all_paragraphs(document):
        if text in paragraph.text:
            return paragraph
    raise ReportDocxError(f"Template slot {text} is unavailable.")


def _render_references(document: Document, report: ReportData) -> None:
    paragraph = _find_paragraph(document, "{{REFERENCES}}")
    references = report["literature_references"]
    if not references:
        _set_paragraph_text(paragraph, "No retained literature references.", size=7)
        return
    current = paragraph
    for index, reference in enumerate(references, start=1):
        if index > 1:
            element = deepcopy(paragraph._p)
            current._p.addnext(element)
            current = Paragraph(element, paragraph._parent)
        title = reference["title"] or reference["source"]
        _set_paragraph_text(current, f"{index}. {title}. ", size=7)
        identifier = f"{reference['identifier_type']}: {reference['identifier']}"
        url = validated_reference_url(reference["canonical_url"])
        if url:
            _add_hyperlink(current, identifier, url, size_half_points=14)
        else:
            run = current.add_run(identifier)
            run.font.name = "Times New Roman"
            run.font.size = Pt(7)


def _table_caption(table) -> str | None:
    caption = table._tbl.tblPr.find(qn("w:tblCaption"))
    return caption.get(qn("w:val")) if caption is not None else None


def _set_table_caption(table, value: str) -> None:
    for name in ("tblCaption", "tblDescription"):
        node = table._tbl.tblPr.find(qn(f"w:{name}"))
        if node is not None:
            node.set(qn("w:val"), value)


def _render_data_sources(document: Document, report: ReportData) -> None:
    table = next(
        (item for item in document.tables if _table_caption(item) == "{{DATA_SOURCES}}"),
        None,
    )
    if table is None:
        raise ReportDocxError("Template Data Sources table is unavailable.")
    _set_table_caption(table, "Data Sources")
    sources = report["data_sources"] or [None]
    prototype = deepcopy(table.rows[1]._tr)
    while len(table.rows) - 1 < len(sources):
        table._tbl.append(deepcopy(prototype))
    while len(table.rows) - 1 > len(sources):
        table._tbl.remove(table.rows[-1]._tr)
    for row, source in zip(table.rows[1:], sources, strict=True):
        if source is None:
            values = ["Not available"] * 6
            url = None
        else:
            status = _status_text(source["status"])
            if source["provider_role"] == "fallback":
                status += " (fallback)"
            record = source["record_identifier"] or source["dataset"] or "Not available"
            values = [
                source["source"],
                source["capability"].replace("_", " "),
                status,
                record,
                source["method"].replace("_", " ") if source["method"] else "Not available",
                (
                    "Open record"
                    if source["human_url"]
                    else "Programmatic annotation source"
                    if source["source"] == "MyVariant.info"
                    else "Not available"
                ),
            ]
            url = validated_reference_url(source["human_url"])
        for cell, value in zip(row.cells, values, strict=True):
            _set_paragraph_text(cell.paragraphs[0], value, size=9.5)
        if url:
            link_paragraph = row.cells[5].paragraphs[0]
            _set_paragraph_text(link_paragraph, "", size=9.5)
            _add_hyperlink(link_paragraph, "Open record", url, size_half_points=19)


def _deterministic_package(data: bytes) -> bytes:
    source = BytesIO(data)
    output = BytesIO()
    with ZipFile(source) as input_zip, ZipFile(
        output, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as output_zip:
        for name in sorted(input_zip.namelist()):
            original = input_zip.getinfo(name)
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            output_zip.writestr(info, input_zip.read(name))
    return output.getvalue()


def render_report_data_docx(
    value: object,
    *,
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
) -> bytes:
    """Populate the authoritative template for one validated allele report."""

    try:
        report = validate_report_data(value)
    except ReportDataError as exc:
        raise ReportDocxError("ReportData V4 is invalid.") from exc
    if report["template_version"] != SUPPORTED_TEMPLATE_VERSION:
        raise ReportDocxError("ReportData requests an unsupported DOCX template.")
    path = Path(template_path)
    if not path.is_file():
        raise ReportDocxError("The authoritative DOCX template is unavailable.")
    document = Document(path)
    template_xml = document._element.xml
    missing = sorted(token for token in REQUIRED_PLACEHOLDERS if token not in template_xml)
    if missing:
        raise ReportDocxError("The authoritative DOCX template is incomplete.")

    _render_references(document, report)
    _render_data_sources(document, report)
    main_table = next(
        (
            item
            for item in document.tables
            if _table_caption(item) == "{{MAIN_FINDINGS_TABLE}}"
        ),
        None,
    )
    if main_table is None:
        raise ReportDocxError("Template Main Findings table is unavailable.")
    _set_table_caption(main_table, "Main Findings")
    replacements = _placeholder_map(report)
    for paragraph in _all_paragraphs(document):
        _replace_paragraph_text(paragraph, replacements)

    if PLACEHOLDER_PATTERN.search(document._element.xml):
        raise ReportDocxError("The generated DOCX contains unresolved placeholders.")
    document.core_properties.title = (
        f"Variant {report['input_index'] + 1:03d} Clinical Report"
    )
    document.core_properties.subject = "Editable allele-level decision-support report"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.keywords = ""
    document.core_properties.comments = ""
    generated_at = datetime.fromisoformat(
        report["provenance"]["generated_at"].replace("Z", "+00:00")
    ).replace(tzinfo=None)
    document.core_properties.created = generated_at
    document.core_properties.modified = generated_at

    output = BytesIO()
    document.save(output)
    data = _deterministic_package(output.getvalue())
    if not data.startswith(b"PK") or len(data) > MAX_DOCX_BYTES:
        raise ReportDocxError("The generated DOCX is invalid or too large.")
    return data


def write_report_data_docx(
    value: object,
    output_dir: str | Path,
    *,
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
) -> Path:
    """Write one independent per-variant DOCX artifact through production rendering."""

    filename = report_docx_filename(value)
    data = render_report_data_docx(value, template_path=template_path)
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / filename
    destination.write_bytes(data)
    return destination


__all__ = [
    "DEFAULT_TEMPLATE_PATH",
    "MAX_DOCX_BYTES",
    "ReportDocxError",
    "SUPPORTED_TEMPLATE_VERSION",
    "render_report_data_docx",
    "report_docx_filename",
    "write_report_data_docx",
]
