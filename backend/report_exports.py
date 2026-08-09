"""Bounded, offline PDF and Word exports for clinical text reports."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from backend.references import validated_reference_url


MAX_REPORT_EXPORT_INPUT_BYTES = 4 * 1024 * 1024
MAX_REPORT_EXPORT_OUTPUT_BYTES = 5 * 1024 * 1024
DEFAULT_REPORT_TITLE = "Clinical Variant Interpretation Report"

_HEADING_PATTERN = re.compile(r"^(#{1,3})\s+(.+?)\s*$")
_HEADING_UNDERLINE_PATTERN = re.compile(r"^([=\-~])\1{2,}$")
_BULLET_PATTERN = re.compile(r"^\s*[-*+]\s+(.+?)\s*$")
_CONTROL_PATTERN = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MARKDOWN_ESCAPE_PATTERN = re.compile(r"\\([\\`*_{}\[\]()#+\-.!>])")
_MARKDOWN_LINK_PATTERN = re.compile(r"\[([^\]]+)\]\((https://[^\s]+)\)")

_NAVY = colors.HexColor("#17324D")
_BLUE = colors.HexColor("#2E74B5")
_LIGHT_BLUE = colors.HexColor("#EAF2F8")
_GRAY = colors.HexColor("#5B6770")


class ReportExportError(ValueError):
    """Raised when report content cannot be exported safely."""


@dataclass(frozen=True, slots=True)
class ReportBlock:
    """One normalized structural block from the generated report."""

    kind: str
    text: str
    level: int = 0
    raw_text: str | None = None


def _normalize_report_text(report_text: str) -> str:
    if not isinstance(report_text, str):
        raise ReportExportError("Report content must be text.")
    normalized = (
        report_text.replace("\r\n", "\n").replace("\r", "\n").strip()
    )
    if not normalized:
        raise ReportExportError("Report content is empty.")
    if len(normalized.encode("utf-8")) > MAX_REPORT_EXPORT_INPUT_BYTES:
        raise ReportExportError("Report content exceeds the export size limit.")
    if _CONTROL_PATTERN.search(normalized):
        raise ReportExportError("Report content contains unsupported characters.")
    return normalized


def _plain_text(value: str) -> str:
    value = _MARKDOWN_LINK_PATTERN.sub(r"\1", value)
    value = _MARKDOWN_ESCAPE_PATTERN.sub(r"\1", value.strip())
    return (
        value.replace("\u00a0", " ")
        .replace("\u2011", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
    )


def _parse_report_blocks(report_text: str) -> list[ReportBlock]:
    blocks: list[ReportBlock] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(
                ReportBlock(
                    kind="paragraph",
                    text=_plain_text(" ".join(paragraph_lines)),
                    raw_text=" ".join(paragraph_lines),
                )
            )
            paragraph_lines.clear()

    lines = report_text.splitlines()
    line_index = 0
    while line_index < len(lines):
        raw_line = lines[line_index]
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            line_index += 1
            continue
        if line_index + 1 < len(lines):
            underline = _HEADING_UNDERLINE_PATTERN.fullmatch(
                lines[line_index + 1].strip()
            )
            if underline:
                flush_paragraph()
                level = {"=": 1, "-": 2, "~": 3}[
                    underline.group(1)
                ]
                blocks.append(
                    ReportBlock(
                        kind="heading",
                        text=_plain_text(line),
                        level=level,
                        raw_text=line,
                    )
                )
                line_index += 2
                continue
        heading = _HEADING_PATTERN.match(line)
        if heading:
            flush_paragraph()
            blocks.append(
                ReportBlock(
                    kind="heading",
                    text=_plain_text(heading.group(2)),
                    level=len(heading.group(1)),
                    raw_text=heading.group(2),
                )
            )
            line_index += 1
            continue
        bullet = _BULLET_PATTERN.match(line)
        if bullet:
            flush_paragraph()
            blocks.append(
                ReportBlock(
                    kind="bullet",
                    text=_plain_text(bullet.group(1)),
                    raw_text=bullet.group(1),
                )
            )
            line_index += 1
            continue
        paragraph_lines.append(line)
        line_index += 1

    flush_paragraph()
    return blocks


def _pdf_text(value: str) -> str:
    return html.escape(value, quote=False)


def _pdf_rich_text(value: str) -> str:
    """Convert allowlisted Markdown links to ReportLab hyperlink markup."""

    parts: list[str] = []
    cursor = 0
    for match in _MARKDOWN_LINK_PATTERN.finditer(value):
        parts.append(html.escape(value[cursor:match.start()], quote=False))
        label = _plain_text(match.group(1))
        url = validated_reference_url(match.group(2))
        if url is None:
            parts.append(html.escape(label, quote=False))
        else:
            parts.append(
                f'<link href="{html.escape(url, quote=True)}" '
                'color="#2E74B5"><u>'
                f"{html.escape(label, quote=False)}</u></link>"
            )
        cursor = match.end()
    parts.append(html.escape(value[cursor:], quote=False))
    return "".join(parts)


def _pdf_footer(canvas: object, document: object) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D5DEE7"))
    canvas.setLineWidth(0.5)
    canvas.line(
        document.leftMargin,
        0.55 * inch,
        letter[0] - document.rightMargin,
        0.55 * inch,
    )
    canvas.setFillColor(_GRAY)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(
        document.leftMargin,
        0.36 * inch,
        "Clinical Variant Interpretation Report",
    )
    page_text = f"Page {document.page}"
    canvas.drawRightString(
        letter[0] - document.rightMargin,
        0.36 * inch,
        page_text,
    )
    canvas.restoreState()


def render_report_pdf(
    report_text: str,
    *,
    title: str = DEFAULT_REPORT_TITLE,
) -> bytes:
    """Render validated report text as an in-memory PDF document."""

    normalized = _normalize_report_text(report_text)
    blocks = _parse_report_blocks(normalized)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.68 * inch,
        title=_plain_text(title),
        author="Clinical Variant Interpretation Project",
        subject="Evidence-bound clinical variant interpretation report",
        creator="Clinical Variant Interpretation Project",
        pageCompression=1,
    )
    sample_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ClinicalReportTitle",
        parent=sample_styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=_NAVY,
        alignment=TA_LEFT,
        spaceAfter=14,
    )
    heading_styles = {
        1: title_style,
        2: ParagraphStyle(
            "ClinicalReportHeading1",
            parent=sample_styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            textColor=_BLUE,
            spaceBefore=11,
            spaceAfter=5,
            keepWithNext=False,
        ),
        3: ParagraphStyle(
            "ClinicalReportHeading2",
            parent=sample_styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11.5,
            leading=14,
            textColor=_NAVY,
            spaceBefore=7,
            spaceAfter=3,
            keepWithNext=False,
        ),
    }
    body_style = ParagraphStyle(
        "ClinicalReportBody",
        parent=sample_styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=12,
        textColor=_NAVY,
        alignment=TA_LEFT,
        spaceAfter=4,
        splitLongWords=True,
    )
    bullet_style = ParagraphStyle(
        "ClinicalReportBullet",
        parent=body_style,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=2,
    )
    notice_style = ParagraphStyle(
        "ClinicalReportNotice",
        parent=body_style,
        backColor=_LIGHT_BLUE,
        borderColor=colors.HexColor("#BED2E3"),
        borderWidth=0.6,
        borderPadding=8,
        leading=12.5,
        spaceBefore=7,
        spaceAfter=7,
    )

    story: list[object] = []
    previous_heading = ""
    for block in blocks:
        if block.kind == "bullet":
            story.append(
                Paragraph(
                    '<font color="#2E74B5">&#8226;</font>&nbsp;&nbsp;'
                    f"{_pdf_rich_text(block.raw_text or block.text)}",
                    bullet_style,
                )
            )
            continue
        if block.kind == "heading":
            previous_heading = block.text
            style = heading_styles.get(block.level, heading_styles[3])
            story.append(
                Paragraph(_pdf_rich_text(block.raw_text or block.text), style)
            )
        else:
            style = (
                notice_style
                if previous_heading.casefold() == "medical disclaimer"
                else body_style
            )
            story.append(
                Paragraph(_pdf_rich_text(block.raw_text or block.text), style)
            )

    if not any(block.kind == "heading" and block.level == 1 for block in blocks):
        story.insert(0, Paragraph(_pdf_text(_plain_text(title)), title_style))
    story.append(Spacer(1, 4))

    document.build(
        story,
        onFirstPage=_pdf_footer,
        onLaterPages=_pdf_footer,
    )
    data = output.getvalue()
    if not data.startswith(b"%PDF-") or len(data) > MAX_REPORT_EXPORT_OUTPUT_BYTES:
        raise ReportExportError("PDF export could not be created safely.")
    return data


def _set_cell_free_page_number(paragraph: object) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(31, 50, 71)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, RGBColor(46, 116, 181), 18, 10),
        "Heading 2": (13, RGBColor(46, 116, 181), 14, 7),
        "Heading 3": (12, RGBColor(31, 77, 120), 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.font.color.rgb = RGBColor(31, 50, 71)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "CLINICAL VARIANT INTERPRETATION REPORT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(91, 103, 112)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _set_cell_free_page_number(footer)
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(91, 103, 112)


def _add_docx_hyperlink(paragraph: object, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_docx_rich_text(paragraph: object, value: str) -> None:
    cursor = 0
    for match in _MARKDOWN_LINK_PATTERN.finditer(value):
        prefix = _plain_text(value[cursor:match.start()])
        if prefix:
            paragraph.add_run(prefix)
        label = _plain_text(match.group(1))
        url = validated_reference_url(match.group(2))
        if url is None:
            paragraph.add_run(label)
        else:
            _add_docx_hyperlink(paragraph, label, url)
        cursor = match.end()
    suffix = _plain_text(value[cursor:])
    if suffix:
        paragraph.add_run(suffix)


def render_report_docx(
    report_text: str,
    *,
    title: str = DEFAULT_REPORT_TITLE,
) -> bytes:
    """Render validated report text as an in-memory Word document."""

    normalized = _normalize_report_text(report_text)
    blocks = _parse_report_blocks(normalized)
    document = Document()
    _configure_docx(document)
    document.core_properties.title = _plain_text(title)
    document.core_properties.subject = (
        "Evidence-bound clinical variant interpretation report"
    )
    document.core_properties.author = (
        "Clinical Variant Interpretation Project"
    )
    document.core_properties.keywords = (
        "clinical variant, evidence report, decision support"
    )

    has_primary_title = any(
        block.kind == "heading" and block.level == 1 for block in blocks
    )
    if not has_primary_title:
        title_paragraph = document.add_paragraph()
        title_paragraph.style = document.styles["Heading 1"]
        title_paragraph.add_run(_plain_text(title))

    previous_heading = ""
    for block in blocks:
        if block.kind == "heading":
            previous_heading = block.text
            paragraph = document.add_paragraph()
            paragraph.style = document.styles[
                f"Heading {min(block.level, 3)}"
            ]
            _add_docx_rich_text(paragraph, block.raw_text or block.text)
            continue
        if block.kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            _add_docx_rich_text(paragraph, block.raw_text or block.text)
            continue
        paragraph = document.add_paragraph()
        _add_docx_rich_text(paragraph, block.raw_text or block.text)
        if previous_heading.casefold() == "medical disclaimer":
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EAF2F8")
            paragraph._p.get_or_add_pPr().append(shading)

    output = BytesIO()
    document.save(output)
    data = output.getvalue()
    if not data.startswith(b"PK") or len(data) > MAX_REPORT_EXPORT_OUTPUT_BYTES:
        raise ReportExportError("Word export could not be created safely.")
    return data


__all__ = [
    "DEFAULT_REPORT_TITLE",
    "MAX_REPORT_EXPORT_INPUT_BYTES",
    "MAX_REPORT_EXPORT_OUTPUT_BYTES",
    "ReportExportError",
    "render_report_docx",
    "render_report_pdf",
]
