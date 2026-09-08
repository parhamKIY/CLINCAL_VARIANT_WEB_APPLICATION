"""Safe Excel source-table parsing for explicitly selected variant rows."""

from __future__ import annotations

from io import BytesIO
from math import isfinite
from typing import Final
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
import io
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.vcf_processing import (
    VCFProcessingError,
    VariantData,
    parse_manual_variants,
)
from config import MAX_VARIANTS_PER_ANALYSIS


SUPPORTED_EXCEL_SUFFIXES: Final = (".xlsx",)
REQUIRED_EXCEL_COLUMNS: Final = ("chrom", "pos", "ref", "alt")
OPTIONAL_EXCEL_COLUMNS: Final = ("qual", "filter")
MAX_EXCEL_ROWS_SCANNED: Final = 1_000
EXCEL_COLUMN_ALIASES: Final = {
    "#chrom": "chrom",
    "chr": "chrom",
    "chrom": "chrom",
    "chromosome": "chrom",
    "pos": "pos",
    "position": "pos",
    "start": "pos",
    "end": "end",
    "ref": "ref",
    "reference": "ref",
    "alt": "alt",
    "alternate": "alt",
    "alternative": "alt",
    "qual": "qual",
    "quality": "qual",
    "filter": "filter",
    "filter_status": "filter",
    "depth": "depth",
    "dp": "depth",
    "ad": "ad",
    "gt_quality": "gq",
    "gq": "gq",
}


class ExcelProcessingError(VCFProcessingError):
    """Raised when an Excel variant table cannot be processed safely."""


def _normalized_header(value: object) -> str | None:
    """Return one normalized header label or None for an unused column."""

    if not isinstance(value, str):
        return None
    normalized = "_".join(value.strip().casefold().split())
    return normalized or None


def _column_indexes(header: tuple[object, ...]) -> dict[str, int]:
    """Map deterministic Excel aliases to canonical manual-table fields."""

    indexes: dict[str, int] = {}
    for index, raw_header in enumerate(header):
        normalized = _normalized_header(raw_header)
        canonical = EXCEL_COLUMN_ALIASES.get(normalized or "")
        if canonical is None:
            continue
        if canonical in indexes:
            raise ExcelProcessingError(
                f"Excel worksheet 1 contains duplicate {canonical.upper()} "
                "columns."
            )
        indexes[canonical] = index

    missing = [
        column
        for column in REQUIRED_EXCEL_COLUMNS
        if column not in indexes
    ]
    if missing:
        labels = ", ".join(column.upper() for column in missing)
        raise ExcelProcessingError(
            f"Excel worksheet 1 is missing required columns: {labels}."
        )
    return indexes


def _is_blank(value: object) -> bool:
    """Return whether an Excel cell has no submitted value."""

    return value is None or (
        isinstance(value, str) and not value.strip()
    )


def _excel_position(value: object, row_number: int) -> object:
    """Normalize an Excel integer without weakening backend validation."""

    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float) and isfinite(value) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        normalized = value.strip()
        if normalized.isdecimal():
            return int(normalized)
    raise ExcelProcessingError(
        f"Excel worksheet 1 row {row_number} POS must be an integer."
    )


def _excel_alternate(value: object) -> object:
    """Retain source ALT; ANNOVAR-like zero is not a canonical allele."""

    return value


def _manual_row_from_excel(
    values: tuple[object, ...],
    indexes: dict[str, int],
    row_number: int,
) -> dict[str, object]:
    """Project one Excel row onto the established manual-input contract."""

    def value_for(column: str) -> object:
        index = indexes.get(column)
        return None if index is None or index >= len(values) else values[index]

    chromosome = value_for("chrom")
    if isinstance(chromosome, (int, float)) and not isinstance(
        chromosome,
        bool,
    ):
        chromosome = str(chromosome).removesuffix(".0")

    return {
        "chrom": chromosome,
        "pos": _excel_position(value_for("pos"), row_number),
        "ref": value_for("ref"),
        "alt": _excel_alternate(value_for("alt")),
        "qual": value_for("qual"),
        "filter": value_for("filter"),
    }


_SPREADSHEETML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _is_invalid_numeric_cell_text(text: str | None) -> bool:
    """Return whether a cell typed as numeric contains non-numeric text.

    VCF-exported Excel files (e.g. from ANNOVAR) commonly use '.' or '-'
    for missing values or gap alleles. When written as <c t="n"><v>.</v></c>
    or <c t="n"><v>-</v></c>, openpyxl's _cast_number calls float/int and
    raises ValueError.
    """
    if text is None:
        return False
    stripped = text.strip()
    if not stripped:
        return False
    if stripped in {".", "-", "-.", "NA", "N/A", "null", "None"}:
        return True
    try:
        if "." in stripped or "E" in stripped or "e" in stripped:
            float(stripped)
        else:
            int(stripped)
        return False
    except ValueError:
        return True


def _sanitize_dot_numeric_cells(payload: bytes) -> bytes:
    """Rewrite xlsx XML to fix cells where non-numeric text is mistyped as numeric.

    VCF-exported Excel files commonly use '.' for missing values and '-' for
    gap INDEL alleles. When the export tool writes these into the XML as
    ``<c t="n"><v>.</v></c>`` or ``<c t="n"><v>-</v></c>``, openpyxl's
    ``_cast_number`` calls ``float('.')`` or ``int('-')`` and raises ValueError.

    This pre-processor rewrites such cells to inline-string type so openpyxl
    sees them as text instead of crashing.
    """

    try:
        zin = zipfile.ZipFile(io.BytesIO(payload), "r")
    except (BadZipFile, OSError):
        return payload  # not a valid zip; let openpyxl report the real error

    needs_rewrite = False
    ns = _SPREADSHEETML_NS
    ET.register_namespace("", ns)

    try:
        for name in zin.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            tree = ET.parse(io.BytesIO(zin.read(name)))
            for c in tree.getroot().iter(f"{{{ns}}}c"):
                if c.get("t", "n") != "n":
                    continue
                v = c.find(f"{{{ns}}}v")
                if v is not None and _is_invalid_numeric_cell_text(v.text):
                    needs_rewrite = True
                    break
            if needs_rewrite:
                break
    except (ET.ParseError, KeyError, OSError):
        zin.close()
        return payload

    if not needs_rewrite:
        zin.close()
        return payload

    buf = io.BytesIO()
    try:
        with zipfile.ZipFile(buf, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if (
                    item.filename.startswith("xl/worksheets/")
                    and item.filename.endswith(".xml")
                ):
                    tree = ET.parse(io.BytesIO(data))
                    root = tree.getroot()
                    for c in root.iter(f"{{{ns}}}c"):
                        if c.get("t", "n") != "n":
                            continue
                        v = c.find(f"{{{ns}}}v")
                        if v is not None and _is_invalid_numeric_cell_text(v.text):
                            c.set("t", "inlineStr")
                            cell_text = v.text
                            c.remove(v)
                            is_el = ET.SubElement(c, f"{{{ns}}}is")
                            t_el = ET.SubElement(is_el, f"{{{ns}}}t")
                            t_el.text = cell_text
                    out = io.BytesIO()
                    tree.write(out, xml_declaration=True, encoding="UTF-8")
                    data = out.getvalue()
                zout.writestr(item, data)
    finally:
        zin.close()

    return buf.getvalue()


def parse_excel_variants(payload: bytes) -> list[VariantData]:
    """Read only worksheet 1 and return the shared normalized variant form."""

    if not isinstance(payload, bytes) or not payload:
        raise ExcelProcessingError("The Excel upload is empty or invalid.")

    payload = _sanitize_dot_numeric_cells(payload)

    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        if not workbook.sheetnames:
            raise ExcelProcessingError(
                "The Excel workbook does not contain a worksheet."
            )

        # Access exactly the first worksheet. No later worksheet is iterated.
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = tuple(next(rows))
        except StopIteration as exc:
            raise ExcelProcessingError(
                "Excel worksheet 1 is empty."
            ) from exc

        indexes = _column_indexes(header)
        recognized_indexes = tuple(indexes.values())
        manual_rows: list[dict[str, object]] = []
        for scanned_count, raw_values in enumerate(rows, start=1):
            row_number = scanned_count + 1
            if scanned_count > MAX_EXCEL_ROWS_SCANNED:
                raise ExcelProcessingError(
                    "Excel worksheet 1 exceeds the safe row-scan limit."
                )
            values = tuple(raw_values)
            if all(
                _is_blank(values[index] if index < len(values) else None)
                for index in recognized_indexes
            ):
                continue
            manual_rows.append(
                _manual_row_from_excel(values, indexes, row_number)
            )
            if len(manual_rows) > MAX_VARIANTS_PER_ANALYSIS:
                raise ExcelProcessingError(
                    "Excel worksheet 1 cannot contain more than "
                    f"{MAX_VARIANTS_PER_ANALYSIS} variant rows."
                )

        if not manual_rows:
            raise ExcelProcessingError(
                "Excel worksheet 1 must contain at least one variant row."
            )
        try:
            return parse_manual_variants(manual_rows)
        except VCFProcessingError as exc:
            raise ExcelProcessingError(str(exc)) from exc
    except ExcelProcessingError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        ParseError,
        ValueError,
    ) as exc:
        raise ExcelProcessingError(
            "The .xlsx upload is damaged or invalid."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _worksheet_for_name(
    workbook: object,
    worksheet_name: str | None,
) -> tuple[object, str]:
    """Return one named worksheet, retaining legacy sheet-1 parsing only."""

    sheetnames = getattr(workbook, "sheetnames", ())
    if not sheetnames:
        raise ExcelProcessingError("The Excel workbook does not contain a worksheet.")
    if worksheet_name is None:
        name = sheetnames[0]
        return workbook[name], "Excel worksheet 1"
    if not isinstance(worksheet_name, str) or worksheet_name not in sheetnames:
        raise ExcelProcessingError(
            "The Excel workbook does not contain worksheet "
            f"{worksheet_name!r}."
        )
    return workbook[worksheet_name], f"Excel worksheet {worksheet_name!r}"


def parse_excel_input_records(
    payload: bytes,
    *,
    worksheet_name: str | None = None,
) -> list[dict[str, object]]:
    """Read one source worksheet without converting ANNOVAR-like alleles.

    ``worksheet_name`` is required by the Stage-3 UI.  Omitting it retains the
    established sheet-1 API for durable pre-Stage-3 callers only.
    """

    if not isinstance(payload, bytes) or not payload:
        raise ExcelProcessingError("The Excel upload is empty or invalid.")
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(_sanitize_dot_numeric_cells(payload)),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        worksheet, worksheet_label = _worksheet_for_name(
            workbook,
            worksheet_name,
        )
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = tuple(next(rows))
        except StopIteration as exc:
            raise ExcelProcessingError(f"{worksheet_label} is empty.") from exc
        try:
            indexes = _column_indexes(header)
        except ExcelProcessingError as exc:
            raise ExcelProcessingError(
                str(exc).replace("Excel worksheet 1", worksheet_label)
            ) from exc
        recognized_indexes = tuple(indexes.values())
        records: list[dict[str, object]] = []
        for scanned_count, raw_values in enumerate(rows, start=1):
            row_number = scanned_count + 1
            if scanned_count > MAX_EXCEL_ROWS_SCANNED:
                raise ExcelProcessingError(
                    f"{worksheet_label} exceeds the safe row-scan limit."
                )
            values = tuple(raw_values)
            if all(_is_blank(values[index] if index < len(values) else None) for index in recognized_indexes):
                continue
            def value_for(column: str) -> object:
                index = indexes.get(column)
                return None if index is None or index >= len(values) else values[index]
            chromosome = value_for("chrom")
            if isinstance(chromosome, (int, float)) and not isinstance(chromosome, bool):
                chromosome = str(chromosome).removesuffix(".0")
            records.append(
                {
                    "worksheet": worksheet.title,
                    "row": row_number,
                    "chrom": chromosome,
                    "start": _excel_position(value_for("pos"), row_number),
                    "end": (
                        _excel_position(value_for("end"), row_number)
                        if value_for("end") is not None
                        else None
                    ),
                    "ref": value_for("ref"),
                    "alt": value_for("alt"),
                    "qual": value_for("qual"),
                    "filter": value_for("filter"),
                    "depth": value_for("depth"),
                    "ad": value_for("ad"),
                    "gq": value_for("gq"),
                }
            )
        if not records:
            raise ExcelProcessingError(
                f"{worksheet_label} must contain at least one variant row."
            )
        return records
    except ExcelProcessingError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
        raise ExcelProcessingError("The .xlsx upload is damaged or invalid.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def discover_excel_worksheets(payload: bytes) -> list[dict[str, object]]:
    """Return actual selectable worksheets and their unranked source-row counts."""

    if not isinstance(payload, bytes) or not payload:
        raise ExcelProcessingError("The Excel upload is empty or invalid.")
    workbook = None
    try:
        sanitized = _sanitize_dot_numeric_cells(payload)
        workbook = load_workbook(
            BytesIO(sanitized),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        names = tuple(workbook.sheetnames)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
        raise ExcelProcessingError("The .xlsx upload is damaged or invalid.") from exc
    finally:
        if workbook is not None:
            workbook.close()

    discovered: list[dict[str, object]] = []
    for name in names:
        records = parse_excel_input_records(payload, worksheet_name=name)
        discovered.append({"name": name, "candidate_count": len(records)})
    return discovered


__all__ = [
    "EXCEL_COLUMN_ALIASES",
    "ExcelProcessingError",
    "OPTIONAL_EXCEL_COLUMNS",
    "REQUIRED_EXCEL_COLUMNS",
    "SUPPORTED_EXCEL_SUFFIXES",
    "discover_excel_worksheets",
    "parse_excel_input_records",
    "parse_excel_variants",
]
