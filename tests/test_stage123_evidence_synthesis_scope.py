"""Stage 123 acceptance: retain the source-attributed evidence-synthesis boundary."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from backend.final_clinical_report import FINAL_CLINICAL_REPORT_SCOPE_NOTICE
from backend.report_docx import render_report_data_docx


pytestmark = pytest.mark.stage123_evidence_synthesis_scope
ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "stage83_golden_report_data_v4.json"


def test_docx_heading_and_result_preserve_source_only_classification_context() -> None:
    document = Document(BytesIO(render_report_data_docx(json.loads(FIXTURE.read_text()))))
    headings = [paragraph.text for paragraph in document.paragraphs]
    result_text = document.tables[1].cell(0, 0).text

    assert "Variant and source classification context:" in headings
    assert "Conclusive Result(s):" not in headings
    assert "System classification: Not independently determined" in result_text
    assert "No independent application classification" in result_text


def test_review_and_final_report_language_disclaims_sign_out_and_adjudication() -> None:
    review_ui = (ROOT / "frontend" / "evidence_review.py").read_text(encoding="utf-8")
    preview_ui = (ROOT / "frontend" / "report_preview.py").read_text(encoding="utf-8")
    final_report = (ROOT / "backend" / "final_clinical_report.py").read_text(
        encoding="utf-8"
    )

    assert "confirmation does not call a model. It confirms evidence review" in review_ui
    assert "not a pathogenicity classification or laboratory sign-out" in review_ui
    assert "audited report disposition, not a " in review_ui
    assert "laboratory sign-out or independent ACMG/AMP adjudication" in review_ui
    assert "Variant and source classification context" in preview_ui
    assert "System classification: Not independently determined" in preview_ui
    assert "independently classify variant pathogenicity" in FINAL_CLINICAL_REPORT_SCOPE_NOTICE
    assert "FINAL_CLINICAL_REPORT_SCOPE_NOTICE" in final_report
