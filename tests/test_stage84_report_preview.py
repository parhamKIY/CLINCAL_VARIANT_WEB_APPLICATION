"""Stage 84 report-first HTML preview acceptance tests."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
import streamlit as st
from docx import Document
from streamlit.testing.v1 import AppTest

from backend.pipeline import create_pipeline_result
from backend.report_data_projection import build_report_data_from_draft
from backend.report_docx import render_report_data_docx
import frontend.ui as frontend_ui
from frontend.report_preview import (
    render_draft_report_preview_pages,
    stable_allele_identity,
)


pytestmark = pytest.mark.stage84_report_preview
PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _draft_report() -> dict[str, object]:
    content = {
        "variant_summary": {
            "display_label": "GENE<1> — GRCh38 2:100001 C>T",
            "assembly": "GRCh38",
            "chrom": "2",
            "pos": 100001,
            "ref": "C",
            "alt": "T",
            "gene": "GENE<1>",
            "gene_id": "HGNC:1",
            "transcript": "NM_000001.1",
            "hgvs_c": "c.101C>T",
            "hgvs_p": "p.(Arg34Trp)",
            "consequence": "missense_variant",
            "rsid": "rs123456",
        },
        "phenotype_context": {
            "accepted_hpo_terms": ["HP:0001250"],
            "matched_hpo_terms": ["HP:0001250"],
            "phenotype_score": 0.68,
            "phenotype_status": "partially_supported",
            "phenotype_to_gene_summary": ["Source-supported gene context."],
            "disease_context": ["Synthetic condition (MONDO:0000001)."],
        },
        "call_quality": {
            "schema_version": "1.0",
            "qual": 50.0,
            "filter": "PASS",
            "status": "passed",
            "acknowledged_at": None,
            "override": None,
        },
        "evidence_sections": [
            {
                "source": "Ensembl VEP",
                "status": "success",
                "items": [
                    {"label": "Consequence", "value": "missense_variant"},
                    {"label": "Coding HGVS", "value": "c.101C>T"},
                ],
            },
            {
                "source": "NCBI ClinVar",
                "status": "success",
                "items": [
                    {"label": "Significance", "value": "uncertain significance"},
                    {
                        "label": "Review status",
                        "value": "criteria provided, single submitter",
                    },
                ],
            },
        ],
        "conflict_summary": {
            "detected": False,
            "status": "none",
            "severity": "none",
            "findings": [],
        },
        "variant_interpretation": {
            "status": "success",
            "narrative": (
                "The exact allele has source-attributed evidence. "
                "The retained evidence remains insufficient for an independent "
                "classification and requires qualified human review."
            ),
            "conflict_assessment": "No material conflict was identified.",
            "warnings": [],
            "model": "synthetic-no-call",
            "prompt_version": "test-v1",
            "generated_at": "2026-08-11T08:00:00Z",
            "failure_type": None,
            "preliminary_classification_status": "classified",
            "preliminary_classification": "Likely pathogenic",
            "classification_rationale": (
                "The retained exact-allele evidence supports a preliminary "
                "Likely pathogenic classification."
            ),
            "limitations": [
                "This is an initial LLM evidence synthesis and requires human review."
            ],
        },
        "reviewer_summary": None,
        "reviewer_notes": [],
        "literature_references": [
            {
                "reference_id": "R1",
                "source": "PubMed",
                "identifier_type": "PMID",
                "identifier": "25741868",
                "title": "Sequence variant interpretation standards",
                "canonical_url": "https://pubmed.ncbi.nlm.nih.gov/25741868/",
                "url_status": "validated",
            }
        ],
        "data_sources": [
            {
                "source": "ClinVar",
                "capability": "clinvar_evidence",
                "status": "available",
                "operational_status": "success",
                "provider_role": "primary",
                "method": "direct_clinvar_lookup",
                "record_identifier": "VCV000000001.1",
                "dataset": None,
                "human_url": None,
                "link_status": "unavailable",
                "fallback_used": False,
                "primary_failure": None,
                "retrieved_at": None,
            }
        ],
        "provenance": {
            "evidence_schema_version": "3.0",
            "interpretation_schema_version": "3.0",
            "providers": ["Ensembl VEP: success", "NCBI ClinVar: success"],
            "upstream_sources": ["Ensembl VEP", "NCBI ClinVar"],
            "interpretation_model": "synthetic-no-call",
            "interpretation_prompt_version": "test-v1",
        },
        "limitations": [
            "Decision-support report only; qualified human review is required."
        ],
    }
    digest = hashlib.sha256(
        json.dumps(
            content,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()[:20]
    return {
        "schema_version": "2.3",
        "report_id": f"dvr-0-{digest}",
        "variant_index": 0,
        "machine_original_report": deepcopy(content),
        "reviewed_report": deepcopy(content),
        "edit_history": [],
        "include_in_final_report": True,
        "selection_history": [],
        "review_status": "draft",
        "created_at": "2026-08-11T08:00:00Z",
        "updated_at": "2026-08-11T08:00:00Z",
    }


def _accept_machine_evidence_change(report: dict[str, object]) -> None:
    content = deepcopy(report["reviewed_report"])
    report["machine_original_report"] = content
    digest = hashlib.sha256(
        json.dumps(
            content,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()[:20]
    report["report_id"] = f"dvr-{report['variant_index']}-{digest}"


def test_preview_has_three_clear_professor_family_pages() -> None:
    pages = render_draft_report_preview_pages(_draft_report())

    assert len(pages) == 3
    assert "NGS Result Report" in pages[0]
    assert "Brief Interpretation(s)" in pages[0]
    assert "Main Finding(s) in Detail" in pages[1]
    assert "Variant interpretation" in pages[1]
    assert "Variant(s) classification" in pages[2]
    assert "Data Sources" in pages[2]
    assert all('class="cv-report-page"' in page for page in pages)


def test_preview_and_docx_show_preliminary_classification_without_a_final_system_call() -> None:
    report = _draft_report()
    pages = render_draft_report_preview_pages(report)
    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'e' * 32}",
    )
    document = Document(BytesIO(render_report_data_docx(report_data)))
    docx_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert "Preliminary evidence-based classification: Likely pathogenic" in pages[0]
    assert "Classification rationale:" in pages[0]
    assert "System classification: Not independently determined" in pages[0]
    assert report_data["preliminary_classification"] == {
        "status": "classified",
        "classification": "Likely pathogenic",
        "rationale": (
            "The retained exact-allele evidence supports a preliminary "
            "Likely pathogenic classification."
        ),
        "limitations": [
            "This is an initial LLM evidence synthesis and requires human review."
        ],
        "review_required": True,
    }
    assert report_data["conclusive_result"]["classification"] is None
    assert "Preliminary evidence-based classification: Likely pathogenic" in docx_text
    assert "Not independently determined" in docx_text


def test_ambiguous_preliminary_output_does_not_force_a_classification() -> None:
    report = _draft_report()
    interpretation = report["machine_original_report"]["variant_interpretation"]
    interpretation.update(
        {
            "preliminary_classification_status": "ambiguous",
            "preliminary_classification": None,
            "classification_rationale": "The retained sources materially disagree.",
            "limitations": ["User review is needed to resolve the conflict."],
        }
    )
    report["reviewed_report"] = deepcopy(report["machine_original_report"])
    _accept_machine_evidence_change(report)

    pages = render_draft_report_preview_pages(report)
    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'f' * 32}",
    )

    assert "Ambiguous — user review required" in pages[0]
    assert "Likely pathogenic" not in pages[0]
    assert report_data["preliminary_classification"]["status"] == "ambiguous"
    assert report_data["preliminary_classification"]["classification"] is None
    assert report_data["conclusive_result"]["classification"] is None


def test_retained_pre_stage125_draft_remains_readable_without_an_invented_label() -> None:
    report = _draft_report()
    report["schema_version"] = "2.2"
    for content_key in ("machine_original_report", "reviewed_report"):
        interpretation = report[content_key]["variant_interpretation"]
        for field in (
            "preliminary_classification_status",
            "preliminary_classification",
            "classification_rationale",
            "limitations",
        ):
            interpretation.pop(field)
    _accept_machine_evidence_change(report)

    pages = render_draft_report_preview_pages(report)
    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'g' * 32}",
    )

    assert "Not available for this pre-Stage-124 analysis" in pages[0]
    assert report_data["preliminary_classification"]["status"] == "unavailable"
    assert report_data["preliminary_classification"]["classification"] is None


def test_preview_uses_assembly_qualified_allele_navigation_identity() -> None:
    report = _draft_report()

    assert stable_allele_identity(report) == "GRCh38 chr2:100001 C>T"
    assert "GRCh38 chr2:100001 C&gt;T" in render_draft_report_preview_pages(
        report
    )[0]


def test_preview_escapes_dynamic_content_and_has_no_executable_script() -> None:
    combined = "".join(render_draft_report_preview_pages(_draft_report()))

    assert "GENE&lt;1&gt;" in combined
    assert "GENE<1>" not in combined
    assert "<script" not in combined.lower()
    assert "javascript:" not in combined.lower()
    assert 'rel="noopener noreferrer"' in combined


def test_preview_exposes_source_findings_without_raw_json() -> None:
    page = render_draft_report_preview_pages(_draft_report())[1]

    assert "Source-attributed main findings" in page
    assert "Ensembl VEP" in page
    assert "NCBI ClinVar" in page
    assert "uncertain significance" in page
    assert "reviewed_user_report" not in page


def test_genebe_only_remains_source_context_not_system_classification() -> None:
    report = _draft_report()
    sections = report["reviewed_report"]["evidence_sections"]
    clinvar = next(section for section in sections if "ClinVar" in section["source"])
    clinvar["items"] = [
        item for item in clinvar["items"] if item["label"] != "Significance"
    ]
    sections.append(
        {
            "source": "GeneBe",
            "status": "success",
            "items": [
                {
                    "label": "Automated ACMG classification",
                    "value": "VUS",
                }
            ],
        }
    )
    _accept_machine_evidence_change(report)

    pages = render_draft_report_preview_pages(report)
    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'c' * 32}",
    )

    assert "System classification: Not independently determined" in pages[0]
    assert "GeneBe automated classification: Uncertain significance (VUS)" in (
        pages[0]
    )
    assert "Direct ClinVar exact classification: not available" in pages[0]
    assert report_data["conclusive_result"] == {
        "gene": "GENE<1>",
        "hgvs_c": "c.101C>T",
        "hgvs_p": "p.(Arg34Trp)",
        "zygosity": None,
        "classification": None,
        "classification_source": None,
        "status": "not_assessed",
    }
    assert report_data["classification_summary"][
        "independent_acmg_adjudication"
    ] is False

    document = Document(BytesIO(render_report_data_docx(report_data)))
    rendered_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert "Not independently determined" in rendered_text
    assert "GeneBe: Uncertain significance (VUS)" in rendered_text


def test_preview_does_not_force_consensus_for_classification_conflict() -> None:
    report = _draft_report()
    report["reviewed_report"]["evidence_sections"].append(
        {
            "source": "GeneBe",
            "status": "success",
            "items": [
                {
                    "label": "Automated ACMG classification",
                    "value": "Likely Pathogenic",
                }
            ],
        }
    )
    _accept_machine_evidence_change(report)

    pages = render_draft_report_preview_pages(report)

    assert "Classification conflict" in pages[0]
    assert "ClinVar: Uncertain significance" in pages[0]
    assert "GeneBe automated: Likely Pathogenic" in pages[0]
    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'a' * 32}",
    )
    assert report_data["conclusive_result"]["classification"] is None
    assert report_data["conclusive_result"]["status"] == "not_assessed"
    assert [
        item["classification"]
        for item in report_data["main_findings"]["classifications"]
    ] == ["uncertain significance", "Likely Pathogenic"]


def test_projection_retains_myvariant_clinvar_derived_rescue() -> None:
    report = _draft_report()
    sections = report["reviewed_report"]["evidence_sections"]
    clinvar = next(section for section in sections if "ClinVar" in section["source"])
    clinvar["status"] = "no_match"
    clinvar["items"] = [
        item for item in clinvar["items"] if item["label"] != "Significance"
    ]
    sections.append(
        {
            "source": "Classification evidence audit",
            "status": "SECONDARY_CLASSIFICATION_EVIDENCE_AVAILABLE",
            "items": [
                {
                    "label": "State",
                    "value": "SECONDARY_CLASSIFICATION_EVIDENCE_AVAILABLE",
                },
                {
                    "label": "MyVariant ClinVar-derived classification",
                    "value": "Pathogenic",
                },
            ],
        }
    )
    _accept_machine_evidence_change(report)

    report_data = build_report_data_from_draft(
        report,
        analysis_id=f"analysis-{'b' * 32}",
    )

    assert report_data["conclusive_result"] == {
        "gene": "GENE<1>",
        "hgvs_c": "c.101C>T",
        "hgvs_p": "p.(Arg34Trp)",
        "zygosity": None,
        "classification": None,
        "classification_source": None,
        "status": "not_assessed",
    }
    assert report_data["main_findings"]["classifications"][-1][
        "independent_evidence"
    ] is False


def test_direct_clinvar_remains_source_context_not_system_classification() -> None:
    report_data = build_report_data_from_draft(
        _draft_report(),
        analysis_id=f"analysis-{'d' * 32}",
    )

    assert report_data["conclusive_result"]["classification"] is None
    assert report_data["conclusive_result"]["classification_source"] is None
    assert report_data["conclusive_result"]["status"] == "not_assessed"
    assert report_data["classification_summary"]["clinvar_classification"] == (
        "uncertain significance"
    )
    assert report_data["main_findings"]["classifications"][0]["source"] == (
        "NCBI ClinVar"
    )


def test_completed_analysis_opens_report_before_provider_details(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = create_pipeline_result()
    result["status"] = "success"
    result["current_stage"] = "completed"
    result["progress_percent"] = 100
    monkeypatch.setattr(
        frontend_ui,
        "render_evidence_review",
        lambda *_args, **_kwargs: st.markdown("STAGE84_REPORT_SURFACE"),
    )
    monkeypatch.setattr(
        frontend_ui,
        "render_analysis_results",
        lambda *_args, **_kwargs: st.markdown("STAGE84_TECHNICAL_SURFACE"),
    )

    app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run(timeout=10)
    app.session_state["pipeline_result"] = result
    app.run(timeout=10)

    assert not app.exception
    selector = next(
        control
        for control in app.button_group
        if control.label == "Analysis result view"
    )
    assert selector.options == [
        "Clinical report review",
        "Analysis and provider details",
    ]
    rendered = [item.value for item in app.markdown]
    assert "STAGE84_REPORT_SURFACE" in rendered
    assert "STAGE84_TECHNICAL_SURFACE" not in rendered
