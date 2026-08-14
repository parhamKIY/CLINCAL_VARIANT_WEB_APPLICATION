"""Golden fidelity gate for the professor-family ReportData V4 DOCX path."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.report_data import validate_report_data
from backend.report_docx import (
    ReportDocxError,
    render_report_data_docx,
    report_docx_filename,
)


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "clinical_variant_report_v1.docx"
FIXTURE = ROOT / "tests" / "fixtures" / "stage83_golden_report_data_v4.json"
GOLDEN = ROOT / "tests" / "golden" / "stage83" / "variant_001_report.docx"
CONTRACT = ROOT / "docs" / "stage_83_fidelity_gate.md"


def _fixture() -> dict[str, object]:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


@pytest.mark.stage83_docx_fidelity
def test_golden_fixture_is_rich_valid_synthetic_report_data() -> None:
    report = validate_report_data(_fixture())

    assert report["input_index"] == 0
    assert report["variant_identity"]["gene"] == "GENE1"
    assert len(report["main_findings"]["population_frequencies"]) == 3
    assert len(report["main_findings"]["computational_evidence"]) == 3
    assert report["main_findings"]["disease_associations"]
    assert report["main_findings"]["stable_variant_ids"]
    assert len(report["literature_references"]) == 3
    assert len(report["data_sources"]) == 8
    assert report_docx_filename(report) == "variant_001_report.docx"


@pytest.mark.stage83_docx_fidelity
def test_production_renderer_is_deterministic_and_matches_golden_artifact() -> None:
    report = _fixture()
    first = render_report_data_docx(report)
    second = render_report_data_docx(report)

    assert first == second
    assert first == GOLDEN.read_bytes()
    assert first.startswith(b"PK")


@pytest.mark.stage83_docx_fidelity
def test_generated_document_preserves_template_visual_system() -> None:
    rendered = render_report_data_docx(_fixture())
    document = Document(BytesIO(rendered))
    section = document.sections[0]

    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(
        round(value.inches, 2) == 1.0
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert document.styles["Report Title"].font.size.pt == 18
    assert document.styles["Major Heading"].font.size.pt == 12
    assert document.styles["Report Reference"].font.size.pt == 7
    assert document.sections[0].header.paragraphs[0].text == ""
    assert document.sections[0].footer.paragraphs[0].text == ""

    with ZipFile(TEMPLATE) as template_zip, ZipFile(BytesIO(rendered)) as rendered_zip:
        assert template_zip.read("word/styles.xml") == rendered_zip.read("word/styles.xml")
        assert template_zip.read("word/theme/theme1.xml") == rendered_zip.read(
            "word/theme/theme1.xml"
        )
        assert not any(name.startswith("word/media/") for name in rendered_zip.namelist())
        xml = rendered_zip.read("word/document.xml").decode("utf-8")
    assert "{{" not in xml
    assert xml.count('w:type="page"') == 2
    assert 'w:color="C00000"' in xml
    assert 'w:fill="D9D9D9"' in xml
    assert xml.count("w:hyperlink") >= 16


@pytest.mark.stage83_docx_fidelity
def test_generated_document_has_professor_family_structure_and_density() -> None:
    document = Document(BytesIO(render_report_data_docx(_fixture())))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    expected_headings = [
        "NGS Result Report",
        "Clinical Features:",
        "Variant and source classification context:",
        "Brief Interpretation(s):",
        "Main Finding(s) in Detail:",
        "Variant interpretation:",
        "Variant(s) classification:",
        "Comments / scope:",
        "References:",
        "Data Sources:",
    ]
    positions = [paragraphs.index(heading) for heading in expected_headings]
    assert positions == sorted(positions)
    assert len(document.tables) == 5
    assert [len(table.rows) for table in document.tables] == [3, 1, 3, 2, 9]
    assert "GENE1:c.101C>T" in document.tables[1].cell(0, 0).text
    assert "Not independently determined" in document.tables[1].cell(0, 0).text
    assert "No independent application classification" in (
        document.tables[1].cell(0, 0).text
    )
    assert "gnomAD" in document.tables[2].cell(1, 2).text
    assert "SIFT" in document.tables[3].cell(0, 0).text
    references = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Report Reference"
    ]
    assert len(references) == 3
    assert all(paragraph.text[0].isdigit() for paragraph in references)


@pytest.mark.stage83_docx_fidelity
def test_generated_tables_retain_fixed_dxa_geometry() -> None:
    document = Document(BytesIO(render_report_data_docx(_fixture())))
    for table in document.tables:
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        indent = tbl_pr.find(qn("w:tblInd"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa"
        assert layout is not None and layout.get(qn("w:type")) == "fixed"
        assert indent is not None and indent.get(qn("w:type")) == "dxa"
        grid = [int(item.get(qn("w:w"))) for item in table._tbl.tblGrid.gridCol_lst]
        assert sum(grid) == int(tbl_w.get(qn("w:w")))
        for row in table.rows:
            assert sum(int(cell._tc.tcPr.tcW.get(qn("w:w"))) for cell in row.cells) == sum(grid)


@pytest.mark.stage83_docx_fidelity
def test_renderer_rejects_template_drift_and_wrong_version(tmp_path: Path) -> None:
    report = _fixture()
    report["template_version"] = "unknown-template"
    with pytest.raises(ReportDocxError, match="unsupported DOCX template"):
        render_report_data_docx(report)

    broken_template = tmp_path / "broken.docx"
    document = Document(TEMPLATE)
    document.paragraphs[0].text = "Missing title slot"
    document.save(broken_template)
    report = _fixture()
    with pytest.raises(ReportDocxError, match="template is incomplete"):
        render_report_data_docx(report, template_path=broken_template)


@pytest.mark.stage83_docx_fidelity
def test_stage83_contract_records_visual_gate_truthfully() -> None:
    contract = " ".join(CONTRACT.read_text(encoding="utf-8").split())
    required = {
        "heading hierarchy",
        "result emphasis",
        "table structure",
        "typography",
        "paragraph flow",
        "page breaks",
        "reference formatting",
        "overall density",
        "clinical-document feel",
        "Stage 115 Microsoft Word visual sign-off",
        "Microsoft Word 16.0",
        "25 pages",
        "no blank pages or page-edge overflow",
    }
    assert all(value in contract for value in required)
