"""Acceptance checks for the Stage 82 authoritative DOCX template."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "clinical_variant_report_v1.docx"
CONTRACT = ROOT / "docs" / "stage_82_template_contract.md"


def _xml() -> str:
    with ZipFile(TEMPLATE) as archive:
        return archive.read("word/document.xml").decode("utf-8")


@pytest.mark.stage82_docx_template
def test_template_is_editable_word_native_and_contains_required_placeholders() -> None:
    assert TEMPLATE.is_file()
    with ZipFile(TEMPLATE) as archive:
        names = set(archive.namelist())
    assert not any(name.startswith("word/media/") for name in names)

    xml = _xml()
    required = {
        "{{REPORT_TITLE}}",
        "{{CLINICAL_FEATURES}}",
        "{{VARIANT_DISPLAY}}",
        "{{ZYGOSITY}}",
        "{{CLASSIFICATION}}",
        "{{BRIEF_INTERPRETATION}}",
        "{{MAIN_FINDINGS_TABLE}}",
        "{{VARIANT_INTERPRETATION}}",
        "{{CLASSIFICATION_SUMMARY}}",
        "{{REFERENCES}}",
        "{{DATA_SOURCES}}",
    }
    assert all(placeholder in xml for placeholder in required)


@pytest.mark.stage82_docx_template
def test_template_owns_reference_page_and_typography_contract() -> None:
    document = Document(TEMPLATE)
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(
        round(value.inches, 2) == 1.0
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert document.styles["Report Title"].font.size.pt == 18
    assert document.styles["Major Heading"].font.size.pt == 12
    assert document.styles["Report Reference"].font.size.pt == 7
    assert len(document.inline_shapes) == 0
    assert len(document.sections[0].header.paragraphs[0].text) == 0
    assert len(document.sections[0].footer.paragraphs[0].text) == 0


@pytest.mark.stage82_docx_template
def test_tables_use_fixed_dxa_geometry_and_repeatable_headers() -> None:
    document = Document(TEMPLATE)
    assert len(document.tables) == 5
    for table in document.tables:
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        indent = tbl_pr.find(qn("w:tblInd"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa"
        assert layout is not None and layout.get(qn("w:type")) == "fixed"
        assert indent is not None and indent.get(qn("w:type")) == "dxa"
        assert all(cell._tc.tcPr.tcW.get(qn("w:type")) == "dxa" for row in table.rows for cell in row.cells)

    xml = _xml()
    assert xml.count("w:tblHeader") >= 4
    assert xml.count("w:cantSplit") >= 3
    assert 'w:color="C00000"' in xml
    assert 'w:fill="D9D9D9"' in xml


@pytest.mark.stage82_docx_template
def test_template_avoids_patient_identifiers_and_out_of_scope_directives() -> None:
    document = Document(TEMPLATE)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    ).lower()
    prohibited = {
        "patient name",
        "national id",
        "medical record number",
        "sanger sequencing",
        "genetic counseling",
        "prenatal",
        "pgd",
        "treatment recommendation",
    }
    assert not any(term in text for term in prohibited)
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""


@pytest.mark.stage82_docx_template
def test_contract_declares_per_variant_artifacts_and_missingness_behavior() -> None:
    contract = CONTRACT.read_text(encoding="utf-8")
    assert "variant_{input_index+1:03d}_report.docx" in contract
    assert "variant_001_report.docx" in contract
    assert "one accepted allele" in contract
    assert "No exact record" in contract
    assert "No supported association" in contract
    assert "Stage 83" in contract
