"""Stage 118 acceptance: provider classifications never become system results."""

from __future__ import annotations

import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from backend.report_data import ReportDataError, validate_report_data
from backend.report_docx import ReportDocxError, render_report_data_docx


pytestmark = pytest.mark.stage118_classification_invariant
_FIXTURE = Path(__file__).parent / "fixtures" / "stage83_golden_report_data_v4.json"


def _source_only_report() -> dict[str, object]:
    return json.loads(_FIXTURE.read_text(encoding="utf-8"))


@pytest.mark.parametrize(
    "classification_source",
    ("NCBI ClinVar", "GeneBe", "MyVariant.info (ClinVar-derived rescue)"),
)
def test_provider_observations_require_an_unassessed_system_result(
    classification_source: str,
) -> None:
    report = _source_only_report()
    findings = report["main_findings"]
    assert isinstance(findings, dict)
    classifications = findings["classifications"]
    assert isinstance(classifications, list)
    classifications[0]["source"] = classification_source

    validated = validate_report_data(report)
    assert validated["conclusive_result"] == {
        "gene": "GENE1",
        "hgvs_c": "c.101C>T",
        "hgvs_p": "p.(Arg34Trp)",
        "zygosity": "heterozygous",
        "classification": None,
        "classification_source": None,
        "status": "not_assessed",
    }
    assert validated["main_findings"]["classifications"][0]["source"] == classification_source


def test_unsafe_provider_result_is_rejected_by_validation_and_docx_rendering() -> None:
    unsafe = _source_only_report()
    conclusive = unsafe["conclusive_result"]
    assert isinstance(conclusive, dict)
    conclusive.update(
        {
            "classification": "Uncertain_significance",
            "classification_source": "NCBI ClinVar",
            "status": "available",
        }
    )

    with pytest.raises(ReportDataError, match="cannot populate conclusive_result"):
        validate_report_data(unsafe)
    with pytest.raises(ReportDocxError, match="ReportData V4 is invalid"):
        render_report_data_docx(unsafe)


def test_docx_keeps_source_context_but_states_no_system_classification() -> None:
    document = Document(BytesIO(render_report_data_docx(_source_only_report())))
    text = "\n".join(
        paragraph.text
        for paragraph in [
            *document.paragraphs,
            *[
                paragraph
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ],
        ]
    )

    assert "System classification: Not independently determined" in text
    assert "NCBI ClinVar: Uncertain significance (VUS)" in text
