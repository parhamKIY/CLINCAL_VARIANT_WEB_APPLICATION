"""Stage 85 document-like report editing acceptance tests."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest
from docx import Document

from backend.evidence_confirmation import confirm_evidence_review
from backend.pipeline import update_draft_variant_report
from backend.report_data_projection import build_report_data_from_draft
from backend.report_docx import render_report_data_docx
from backend.variant_report import save_draft_variant_report
from tests.test_pipeline import (
    TestStage40FrontendReviewWorkflow as _Stage40FrontendReviewWorkflow,
)
from tests.test_stage84_report_preview import _draft_report


pytestmark = pytest.mark.stage85_document_editor
PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _editor_app() -> None:
    import streamlit as st

    from frontend.evidence_review import (
        _render_report_editor,
        _request_report_edit,
    )
    from tests.test_stage84_report_preview import _draft_report

    report = _draft_report()
    if st.button(
        "Open document editor",
        on_click=_request_report_edit,
        args=(report["report_id"],),
    ):
        pass
    _render_report_editor(report, {"evidence_review_reports": []})


def test_editor_is_explicit_and_organized_by_document_region() -> None:
    app = AppTest.from_function(_editor_app).run(timeout=10)

    assert not app.exception
    editor = next(
        item for item in app.status if item.label == "Edit clinical report"
    )
    assert editor.proto.expanded is False
    assert [area.label for area in app.text_area] == [
        "Reviewer summary",
        "Variant interpretation narrative",
        "Conflict assessment wording",
        "Reviewer notes (one note per line)",
    ]
    headings = [item.value for item in app.markdown]
    assert "#### Brief interpretation" in headings
    assert "#### Variant interpretation" in headings
    assert "#### Classification summary" in headings
    assert "#### Reviewer notes" in headings
    assert {button.label for button in app.button} >= {
        "Open document editor",
        "Cancel",
        "Reset editable fields",
        "Save report edits",
    }


def test_edit_button_opens_editor_and_cancel_closes_it() -> None:
    app = AppTest.from_function(_editor_app).run(timeout=10)
    next(
        button
        for button in app.button
        if button.label == "Open document editor"
    ).click().run(timeout=10)

    editor = next(
        item for item in app.status if item.label == "Edit clinical report"
    )
    assert editor.proto.expanded is True

    next(
        button for button in app.button if button.label == "Cancel"
    ).click().run(timeout=10)

    editor = next(
        item for item in app.status if item.label == "Edit clinical report"
    )
    assert editor.proto.expanded is False
    assert not app.exception


def test_report_preview_edit_action_opens_document_editor() -> None:
    app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run(timeout=10)
    app.session_state["pipeline_result"] = (
        _Stage40FrontendReviewWorkflow._draft_result()
    )
    app.run(timeout=10)

    editor = next(
        item for item in app.status if item.label == "Edit clinical report"
    )
    assert editor.proto.expanded is False
    next(
        button for button in app.button if button.label == "Edit"
    ).click().run(timeout=10)

    editor = next(
        item for item in app.status if item.label == "Edit clinical report"
    )
    assert editor.proto.expanded is True
    assert not app.exception
    assert any(
        button.label == "Download editable Word report"
        for button in app.get("download_button")
    )


def test_designated_edits_are_audited_without_mutating_generated_evidence() -> None:
    report = _draft_report()
    original = deepcopy(report["machine_original_report"])

    edited = save_draft_variant_report(
        report,
        reviewer_summary="Reviewer-approved brief interpretation.",
        interpretation_narrative="Reviewer-approved variant interpretation.",
        conflict_assessment="Reviewer-approved classification summary.",
        reviewer_notes=["Reviewer note retained in the report."],
        timestamp="2026-08-11T11:00:00Z",
        reviewer_context="stage85_test",
    )

    assert edited["machine_original_report"] == original
    assert edited["reviewed_report"]["evidence_sections"] == original[
        "evidence_sections"
    ]
    assert [item["field_path"] for item in edited["edit_history"]] == [
        "/reviewer_summary",
        "/variant_interpretation/narrative",
        "/variant_interpretation/conflict_assessment",
        "/reviewer_notes",
    ]


def test_edits_project_to_report_data_and_regenerate_word_artifact() -> None:
    report = _draft_report()
    edited = save_draft_variant_report(
        report,
        reviewer_summary="Reviewer-approved brief interpretation.",
        interpretation_narrative="Reviewer-approved variant interpretation.",
        conflict_assessment="Reviewer-approved classification summary.",
        reviewer_notes=["Reviewer note retained in the report."],
        timestamp="2026-08-11T11:00:00Z",
        reviewer_context="stage85_test",
    )

    report_data = build_report_data_from_draft(
        edited,
        analysis_id="analysis-stage85-document-editor",
    )
    docx_data = render_report_data_docx(report_data)
    document = Document(BytesIO(docx_data))
    rendered_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )

    assert report_data["review_state"]["reviewer_summary"] == (
        "Reviewer-approved brief interpretation."
    )
    assert report_data["interpretation"][
        "current_reviewer_interpretation"
    ] == "Reviewer-approved variant interpretation."
    assert report_data["classification_summary"]["summary"] == (
        "Reviewer-approved classification summary."
    )
    assert "Reviewer-approved brief interpretation." in rendered_text
    assert "Reviewer-approved variant interpretation." in rendered_text
    assert "Reviewer-approved classification summary." in rendered_text
    assert docx_data.startswith(b"PK")


def test_material_report_edit_invalidates_prior_confirmation() -> None:
    result = _Stage40FrontendReviewWorkflow._draft_result()
    evidence_report = result["evidence_review_reports"][0]
    package = confirm_evidence_review(evidence_report)
    result["reviewed_evidence_packages"] = [dict(package)]
    report = result["draft_variant_reports"][0]
    content = report["reviewed_report"]
    edited = save_draft_variant_report(
        report,
        reviewer_summary="Material reviewer edit.",
        interpretation_narrative=content["variant_interpretation"]["narrative"],
        conflict_assessment=content["variant_interpretation"][
            "conflict_assessment"
        ],
        reviewer_notes=content["reviewer_notes"],
        timestamp="2026-08-11T11:00:00Z",
        reviewer_context="stage85_test",
    )

    updated = update_draft_variant_report(result, edited)

    assert updated["reviewed_evidence_packages"] == []
    assert updated["workflow_state"] == "awaiting_final_review"
