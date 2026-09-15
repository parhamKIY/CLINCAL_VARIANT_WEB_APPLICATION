"""Deterministic Stage 3 contracts for selected XLSX rows."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook

from backend.excel_processing import (
    ExcelProcessingError,
    discover_excel_worksheets,
    parse_excel_input_records,
)
from backend.pipeline import run_annovar_like_input_processing
from frontend.xlsx_selection import (
    XLSXSelectionError,
    build_preprocessing_summary,
    select_excel_input_records,
)
from frontend.execution import prepare_analysis_recovery_request


pytestmark = pytest.mark.testing_v3_input


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, count in (("Noticable", 5), ("Top_Phen", 12), ("All_Phen_Related", 74)):
        worksheet = workbook.create_sheet(name)
        worksheet.append(("Chr", "Start", "End", "Ref", "Alt", "Quality", "Filter"))
        for index in range(count):
            worksheet.append(
                (
                    "1",
                    1000 + index,
                    1000 + index,
                    "0" if index == 3 else "A",
                    "0" if index == 2 else "G",
                    1.0 if index == 1 else 10.0,
                    "QDfilter;SORfilter" if index == 2 else "PASS",
                )
            )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_discovery_requires_an_explicit_existing_worksheet() -> None:
    payload = _workbook_bytes()

    assert discover_excel_worksheets(payload) == [
        {"name": "Noticable", "candidate_count": 5},
        {"name": "Top_Phen", "candidate_count": 12},
        {"name": "All_Phen_Related", "candidate_count": 74},
    ]
    assert len(parse_excel_input_records(payload, worksheet_name="Top_Phen")) == 12
    with pytest.raises(ExcelProcessingError, match="does not contain worksheet"):
        parse_excel_input_records(payload, worksheet_name="Unknown")


def test_selected_rows_preserve_source_order_and_reject_no_or_over_limit_selection() -> None:
    records = parse_excel_input_records(_workbook_bytes(), worksheet_name="Noticable")

    selected = select_excel_input_records(records, (6, 2, 4))

    assert [record["row"] for record in selected] == [2, 4, 6]
    assert selected[1]["filter"] == "QDfilter;SORfilter"
    assert selected[1]["alt"] == "0"
    assert records[3]["ref"] == "0"
    with pytest.raises(XLSXSelectionError, match="Select at least one"):
        select_excel_input_records(records, ())
    with pytest.raises(XLSXSelectionError, match="maximum of 10"):
        select_excel_input_records(records * 3, tuple(range(2, 13)))


def test_preprocessing_summary_keeps_selected_and_canonical_counts_separate() -> None:
    summary = build_preprocessing_summary(
        [
            {"status": "ACCEPTED_DIRECT", "canonical_variant_index": 0},
            {"status": "NORMALIZED_AND_ACCEPTED", "canonical_variant_index": 1},
            {"status": "IDENTITY_UNRESOLVED", "canonical_variant_index": None},
        ]
    )

    assert summary == {
        "selected_count": 3,
        "canonical_count": 2,
        "unresolved_count": 1,
        "accepted_direct_count": 1,
        "normalized_count": 1,
    }


def test_recovery_uses_only_explicitly_selected_source_rows() -> None:
    payload = _workbook_bytes()
    records = parse_excel_input_records(payload, worksheet_name="Top_Phen")
    selected = select_excel_input_records(records, (2, 4, 6))

    request = prepare_analysis_recovery_request(
        uploaded_vcf=None,
        manual_variants=None,
        excel_input_records=selected,
        phenotypes=["HP:0001250"],
        llm_model="test-model",
        input_type="excel",
    )

    assert request["manual_variants"] == []
    assert [item["row"] for item in request["excel_input_records"]] == [2, 4, 6]
    assert {item["worksheet"] for item in request["excel_input_records"]} == {"Top_Phen"}


def test_large_worksheet_never_implies_a_default_or_first_ten_selection() -> None:
    records = parse_excel_input_records(
        _workbook_bytes(),
        worksheet_name="All_Phen_Related",
    )

    assert len(records) == 74
    assert select_excel_input_records(records, (2, 11, 20)) == [
        records[0], records[9], records[18]
    ]
    assert len(select_excel_input_records(records, tuple(range(2, 12)))) == 10
    with pytest.raises(XLSXSelectionError, match="maximum of 10"):
        select_excel_input_records(records, tuple(range(2, 13)))


def test_existing_stage2_preprocessing_accounts_for_selected_source_rows_only() -> None:
    source_rows = [
        {
            "worksheet": "Top_Phen",
            "row": 9,
            "chrom": "1",
            "start": 100,
            "end": 100,
            "ref": "A",
            "alt": "G",
            "qual": 1.0,
            "filter": "QDfilter",
            "depth": 2,
            "ad": None,
            "gq": 1,
        },
        {
            "worksheet": "Top_Phen",
            "row": 11,
            "chrom": "1",
            "start": 200,
            "end": 201,
            "ref": "AC",
            "alt": 0,
            "qual": 2.0,
            "filter": "SORfilter",
            "depth": 3,
            "ad": None,
            "gq": 2,
        },
        {
            "worksheet": "Top_Phen",
            "row": 12,
            "chrom": "1",
            "start": 300,
            "end": 300,
            "ref": "0",
            "alt": "N",
            "qual": 3.0,
            "filter": "PASS",
            "depth": 4,
            "ad": None,
            "gq": 3,
        },
    ]

    def reference_fetcher(**kwargs: object) -> dict[str, object]:
        start = kwargs["start"]
        end = kwargs["end"]
        return {
            "status": "success",
            "assembly": "GRCh38",
            "chrom": "1",
            "start": start,
            "end": end,
            "sequence": "T" if start == 199 and end == 199 else "AC",
            "source": "ensembl_grch38_sequence",
            "fallback_used": False,
            "failure_reason": None,
            "provider_attempts": [],
        }

    result = run_annovar_like_input_processing(
        source_rows,
        phenotypes=[],
        reference_fetcher=reference_fetcher,
    )

    assert len(result["input_preprocessing_results"]) == 3
    assert result["variant_count"] == 2
    assert [item["status"] for item in result["input_preprocessing_results"]] == [
        "ACCEPTED_DIRECT",
        "NORMALIZED_AND_ACCEPTED",
        "IDENTITY_UNRESOLVED",
    ]
    assert result["input_preprocessing_results"][2]["canonical_variant"] is None
    assert result["input_preprocessing_results"][0]["source_provenance"]["source_filter"] == "QDfilter"


def test_excel_with_mistyped_numeric_dash_and_dot_cells() -> None:
    """Export tools often mistype '-' or '.' as numeric cells; sanitizer must rescue them."""
    import zipfile
    import xml.etree.ElementTree as ET

    wb = Workbook()
    ws = wb.active
    ws.append(("Chr", "Start", "End", "Ref", "Alt", "Quality", "Filter"))
    ws.append(("1", 1000, 1000, "A", "G", 30, "PASS"))
    buf = BytesIO()
    wb.save(buf)

    zin = zipfile.ZipFile(BytesIO(buf.getvalue()), "r")
    out = BytesIO()
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)

    with zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                tree = ET.parse(BytesIO(data))
                root = tree.getroot()
                for c in root.iter(f"{{{ns}}}c"):
                    if c.get("r") == "E2":
                        c.set("t", "n")
                        for child in list(c):
                            c.remove(child)
                        v = ET.SubElement(c, f"{{{ns}}}v")
                        v.text = "-"
                xml_buf = BytesIO()
                tree.write(xml_buf, xml_declaration=True, encoding="UTF-8")
                data = xml_buf.getvalue()
            zout.writestr(item, data)

    payload = out.getvalue()
    worksheets = discover_excel_worksheets(payload)
    assert len(worksheets) == 1
    assert worksheets[0]["candidate_count"] == 1

    records = parse_excel_input_records(payload, worksheet_name="Sheet")
    assert len(records) == 1
    assert records[0]["alt"] == "-"
    assert records[0]["ref"] == "A"
    assert records[0]["start"] == 1000


def test_excel_gap_indel_deletion_and_insertion_normalization() -> None:
    source_rows = [
        {
            "worksheet": "Top_Phen",
            "row": 9,
            "chrom": "3",
            "start": 46709587,
            "end": 46709589,
            "ref": "AAG",
            "alt": "-",
            "qual": 3570.0,
            "filter": "PASS",
            "depth": "75,75",
            "ad": "64,64",
            "gq": 99,
        },
        {
            "worksheet": "All_Phen_Related",
            "row": 106,
            "chrom": "16",
            "start": 3729216,
            "end": 3729216,
            "ref": "-",
            "alt": "C",
            "qual": 5.0,
            "filter": "PASS",
            "depth": 84,
            "ad": 8,
            "gq": 31,
        },
    ]

    def reference_fetcher(**kwargs: object) -> dict[str, object]:
        chrom = str(kwargs["chrom"])
        start = kwargs["start"]
        end = kwargs["end"]
        if chrom == "3" and start == 46709587 and end == 46709589:
            seq = "AAG"
        elif chrom == "3" and start == 46709586 and end == 46709586:
            seq = "G"
        elif chrom == "16" and start == 3729216 and end == 3729216:
            seq = "G"
        else:
            seq = "N"
        return {
            "status": "success",
            "assembly": "GRCh38",
            "chrom": chrom,
            "start": start,
            "end": end,
            "sequence": seq,
            "source": "ensembl_grch38_sequence",
            "fallback_used": False,
            "failure_reason": None,
            "provider_attempts": [],
        }

    result = run_annovar_like_input_processing(
        source_rows,
        phenotypes=[],
        reference_fetcher=reference_fetcher,
    )

    assert result["variant_count"] == 2
    assert [item["status"] for item in result["input_preprocessing_results"]] == [
        "NORMALIZED_AND_ACCEPTED",
        "NORMALIZED_AND_ACCEPTED",
    ]
    # Deletion: 3:46709587-46709589 AAG>- anchors to 3:46709586 GAAG>G
    assert result["variants"][0]["chrom"] == "3"
    assert result["variants"][0]["pos"] == 46709586
    assert result["variants"][0]["ref"] == "GAAG"
    assert result["variants"][0]["alt"] == "G"
    # Multi-sample depth string '75,75' parsed gracefully
    assert result["input_preprocessing_results"][0]["source_provenance"]["source_depth"] == 75.0
    assert result["input_preprocessing_results"][0]["source_provenance"]["source_alt"] == "-"

    # Insertion: 16:3729216 -> anchors to 16:3729216 G>GC
    assert result["variants"][1]["chrom"] == "16"
    assert result["variants"][1]["pos"] == 3729216
    assert result["variants"][1]["ref"] == "G"
    assert result["variants"][1]["alt"] == "GC"
    assert result["input_preprocessing_results"][1]["source_provenance"]["source_ref"] == "-"


def test_excel_unresolved_variant_generates_pipeline_warning() -> None:
    source_rows = [
        {
            "worksheet": "Top_Phen",
            "row": 2,
            "chrom": "22",
            "start": 37974173,
            "end": 37974173,
            "ref": "A",
            "alt": "T",
            "qual": 10.0,
            "filter": "PASS",
            "depth": 99,
            "ad": 24,
            "gq": 99,
        },
        {
            "worksheet": "Top_Phen",
            "row": 5,
            "chrom": "22",
            "start": 100,
            "end": 100,
            "ref": "N",
            "alt": "G",
            "qual": 10.0,
            "filter": "PASS",
            "depth": 50,
            "ad": 10,
            "gq": 99,
        },
    ]

    result = run_annovar_like_input_processing(
        source_rows,
        phenotypes=[],
    )

    assert result["variant_count"] == 1
    assert any(
        "1 selected input row was (row 5) excluded from analysis" in w
        for w in result["warnings"]
    )


def test_excel_deletion_without_end_column_auto_derives_end() -> None:
    source_rows = [
        {
            "worksheet": "Top_Phen",
            "row": 9,
            "chrom": "3",
            "start": 46709587,
            "end": None,  # No end column in workbook
            "ref": "AAG",
            "alt": "-",
            "qual": 10.0,
            "filter": "PASS",
            "depth": 50,
            "ad": 10,
            "gq": 99,
        },
    ]

    def reference_fetcher(**kwargs: object) -> dict[str, object]:
        start = kwargs["start"]
        end = kwargs["end"]
        if start == 46709587 and end == 46709589:
            seq = "AAG"
        elif start == 46709586 and end == 46709586:
            seq = "G"
        else:
            seq = "N"
        return {
            "status": "success",
            "assembly": "GRCh38",
            "chrom": "3",
            "start": start,
            "end": end,
            "sequence": seq,
            "source": "ensembl_grch38_sequence",
            "fallback_used": False,
            "failure_reason": None,
            "provider_attempts": [],
        }

    result = run_annovar_like_input_processing(
        source_rows,
        phenotypes=[],
        reference_fetcher=reference_fetcher,
    )

    assert result["variant_count"] == 1
    assert result["input_preprocessing_results"][0]["status"] == "NORMALIZED_AND_ACCEPTED"
    assert result["variants"][0]["ref"] == "GAAG"
    assert result["variants"][0]["alt"] == "G"
