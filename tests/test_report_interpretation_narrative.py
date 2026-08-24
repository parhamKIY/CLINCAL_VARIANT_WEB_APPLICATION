"""Regression coverage for interpretation narrative report projection."""

from __future__ import annotations

import json
from copy import deepcopy
from io import BytesIO

from docx import Document

from backend.llm import LLMClient, LLMResponse
from backend.report_data_projection import build_report_data_from_draft
from backend.report_docx import render_report_data_docx
from backend.variant_interpretation import (
    interpret_variant,
    validate_variant_interpretation_result,
)
from backend.variant_report import build_draft_variant_report
from frontend.report_preview import render_draft_report_preview_pages
from test_pipeline import FakeLLMAdapter, TestEvidenceObject as EvidenceFactory


_TIMESTAMP = "2026-08-24T20:00:00Z"
_SUBSTANTIVE_NARRATIVE = (
    "The collected exact-allele evidence supports a cautious uncertain-significance "
    "draft assessment because the source-attributed population observations, ClinVar "
    "record, transcript consequence, gene-disease context, and computational annotations "
    "must be considered together without treating unavailable phenotype support as "
    "negative biological evidence or allowing any one provider classification to replace "
    "the synthesis, while the retained provenance and explicit limitations permit a "
    "qualified reviewer to inspect every supporting source, resolve remaining uncertainty, "
    "and decide whether additional case-specific evidence is needed before approving the "
    "draft for clinical use in the established human-review workflow, which remains the "
    "required safety boundary for this educational decision-support application. "
    "The narrative must remain visible in both the concise report summary and the detailed "
    "interpretation section."
)
_PHENOTYPE_PREFIX = "Phenotype conclusion: no supported association found."


def _no_match_evidence() -> dict[str, object]:
    evidence = deepcopy(EvidenceFactory._complete_evidence_object())
    evidence["matched_hpo_terms"] = []
    evidence["phenotype_score"] = 0.0
    relationship = evidence["phenotype_relationship"]
    relationship["matched_patient_hpo_terms"] = []
    relationship["local_phenotype_score"] = 0.0
    relationship["phenotype_status"] = "no_exact_match"
    return evidence


def _partial_field_result() -> tuple[dict[str, object], dict[str, object]]:
    evidence = _no_match_evidence()
    response = LLMResponse(
        content=json.dumps(
            {
                "ai_classification": "Uncertain significance",
                "interpretation": _SUBSTANTIVE_NARRATIVE,
                "conflict_assessment": "No material source conflict was identified.",
                "phenotype_conclusion": "unsupported free-form conclusion",
                "warnings": [],
            }
        ),
        model="report-narrative-regression-model",
        finish_reason="stop",
    )
    result = interpret_variant(
        evidence,
        client=LLMClient(FakeLLMAdapter(response)),
        timestamp=_TIMESTAMP,
    )
    return evidence, result


def _paragraph_after(document: Document, heading: str) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    position = paragraphs.index(heading)
    return next(text for text in paragraphs[position + 1 :] if text.strip())


def test_valid_narrative_survives_partial_validation_and_every_report_boundary() -> None:
    evidence, result = _partial_field_result()

    assert result["status"] == "success"
    assert result["field_validation"]["phenotype_conclusion"] == "invalid"
    assert result["field_validation"]["interpretation"] == "valid"
    assert result["interpretation"] == (
        f"{_PHENOTYPE_PREFIX}\n\n{_SUBSTANTIVE_NARRATIVE}"
    )

    draft = build_draft_variant_report(evidence, result, variant_index=0)
    report_data = build_report_data_from_draft(
        draft,
        analysis_id=f"analysis-{'d' * 32}",
    )
    expected = result["interpretation"]

    assert draft["reviewed_report"]["variant_interpretation"]["narrative"] == expected
    assert report_data["interpretation"]["original_model_interpretation"] == expected
    assert report_data["interpretation"]["current_reviewer_interpretation"] == expected

    preview_pages = render_draft_report_preview_pages(draft)
    assert "The collected exact-allele evidence" in preview_pages[0]
    assert _SUBSTANTIVE_NARRATIVE in preview_pages[1]

    document = Document(BytesIO(render_report_data_docx(report_data)))
    assert "The collected exact-allele evidence" in _paragraph_after(
        document,
        "Brief Interpretation(s):",
    )
    assert _SUBSTANTIVE_NARRATIVE in (
        "\n".join(paragraph.text for paragraph in document.paragraphs)
    )


def test_historical_result_without_field_validation_keeps_narrative_in_report() -> None:
    evidence, current = _partial_field_result()
    historical = dict(current)
    historical["schema_version"] = "1.2"
    historical.pop("phenotype_conclusion")
    historical.pop("field_validation")

    normalized = validate_variant_interpretation_result(
        historical,
        evidence=evidence,
    )
    draft = build_draft_variant_report(evidence, normalized, variant_index=0)

    assert normalized["interpretation"] == current["interpretation"]
    assert normalized["field_validation"]["interpretation"] == "valid"
    assert "The collected exact-allele evidence" in (
        render_draft_report_preview_pages(draft)[0]
    )
